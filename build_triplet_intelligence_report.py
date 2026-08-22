#!/usr/bin/env python3
"""Build a local-only FSSAI baseline and news-triplet intelligence report.

The analysis deliberately ignores the historical news-triplet status field.
All retained Oil, Ghee, and Milk news triplets are analyzed together; model
confidence is the only model-derived quality variable used in rankings.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean, median

import matplotlib.pyplot as plt
import networkx as nx
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
DATE_TAG = "20260822"
OUT = ROOT / f"FSSAI_Triplet_Intelligence_{DATE_TAG}_release"
FIGURES = OUT / "figures"
TABLES = OUT / "tables"

BASELINE_PATH = ROOT / "triplets_verified.csv"
NEWS_PATH = ROOT / "News_Triplets_Oil_Ghee_Milk_20260820.csv"
ARTICLE_PATHS = [
    ROOT / "qwen_oil_detailed_20260729_230158" / "oil_articles.cleaned.jsonl",
    ROOT / "ghee_qwen_events_20260802_101933" / "ghee_articles.cleaned.jsonl",
    ROOT / "milk_parity_final_20260819_173030" / "blind_recall_articles.jsonl",
]

BASE_COLUMNS = [
    "snippet_id",
    "source_id",
    "source_file",
    "source_type",
    "chunk_index",
    "subject",
    "subject_type",
    "subject_id",
    "predicate",
    "object",
    "object_type",
    "object_id",
    "confidence",
    "evidence_span",
]

DIRECT_FOOD_RELATIONS = {"hasadulterant", "issubstitutefor"}
EFFECT_RELATIONS = {"causeseffect", "associatedwith", "wasassociatedwith", "haseffecttype"}
METHOD_RELATIONS = {
    "detectedby",
    "requireskit",
    "producesindicator",
    "isperformedon",
    "isperformedat",
    "isperformedas",
    "identifies",
    "inputto",
    "outputof",
}
ACTION_RELATIONS = {"triggeredaction", "issuedby", "targets", "citesfinding"}

GENERIC_ADULTERANTS = {
    "adulterant",
    "adulterants",
    "unknown adulterant",
    "unspecified adulterant",
    "foreign substance",
    "foreign substances",
    "harmful chemicals",
    "chemicals",
    "chemical adulterants",
    "poison",
    "toxic substance",
    "contaminants",
    "impurities",
    "adulterated ghee",
    "adulterated milk",
    "adulterated oil",
}

FOOD_ALIASES = [
    ("Milk powder", ["milk powder", "skimmed milk powder", "skim milk powder"]),
    ("Mustard oil", ["mustard oil", "sarson oil"]),
    ("Groundnut oil", ["groundnut oil", "peanut oil"]),
    ("Sunflower oil", ["sunflower oil"]),
    ("Soybean oil", ["soybean oil", "soyabean oil", "soya bean oil"]),
    ("Coconut oil", ["coconut oil"]),
    ("Sesame oil", ["sesame oil", "til oil"]),
    ("Rice bran oil", ["rice bran oil"]),
    ("Palm oil", ["palm oil", "palmolein oil", "palm olein"]),
    ("Edible/cooking oil", ["edible oil", "cooking oil", "vegetable oil"]),
    ("Tirupati laddu", ["tirupati laddu", "tirumala laddu"]),
    ("Ghee", ["ghee", "clarified butter"]),
    ("Milk", ["milk"]),
    ("Paneer", ["paneer"]),
    ("Khoya/mawa", ["khoya", "mawa"]),
    ("Curd/yoghurt", ["curd", "yoghurt", "yogurt"]),
    ("Butter", ["butter"]),
    ("Cheese", ["cheese"]),
]

ADULTERANT_ALIASES = [
    ("ethylene glycol", ["ethylene glycol", "monoethylene glycol", "meg"]),
    ("animal fat", ["animal fat", "animal fats", "rendered animal fat", "rendered animal fats"]),
    ("palm oil family", ["palm oil", "palm oils", "palmolein", "palm olein", "palm stearin", "palm kernel oil"]),
    ("vegetable oil/fat", ["vegetable oil", "vegetable oils", "vegetable fat", "vegetable fats"]),
    ("lard/pig fat", ["lard", "pig fat", "pork fat"]),
    ("beef tallow", ["beef tallow", "beef fat"]),
    ("fish oil", ["fish oil"]),
    ("formalin/formaldehyde", ["formalin", "formaldehyde"]),
    ("detergent", ["detergent", "detergents", "detergent powder", "synthetic detergent"]),
    ("starch", ["starch", "added starch"]),
    ("water", ["water", "added water"]),
    ("urea", ["urea", "added urea"]),
    ("argemone oil", ["argemone oil"]),
    ("hydrogen peroxide", ["hydrogen peroxide"]),
    ("vanaspati", ["vanaspati"]),
    ("mineral oil", ["mineral oil"]),
    ("metanil yellow", ["metanil yellow"]),
    ("rhodamine B", ["rhodamine b", "rhodamine-b"]),
    ("lead chromate", ["lead chromate"]),
    ("artificial colour", ["artificial colour", "artificial color", "synthetic colour", "synthetic color"]),
    ("aflatoxin", ["aflatoxin", "aflatoxin m1"]),
    ("antibiotic residue", ["antibiotic", "antibiotics", "antibiotic residue", "antibiotic residues"]),
    ("pesticide residue", ["pesticide", "pesticides", "pesticide residue", "pesticide residues"]),
    ("microplastics", ["microplastic", "microplastics"]),
    ("oxytocin", ["oxytocin"]),
    ("calcium carbide", ["calcium carbide"]),
    ("ethylene oxide", ["ethylene oxide"]),
    ("heavy metals", ["heavy metal", "heavy metals"]),
    ("acetic acid ester", ["acetic acid ester"]),
]

EFFECT_ALIASES = [
    ("renal failure", ["acute renal failure", "renal failure", "kidney failure", "acute kidney injury"]),
    ("kidney damage", ["kidney damage", "renal damage", "kidney injury"]),
    ("multi-organ failure", ["multi-organ failure", "multiple organ failure", "multi organ failure", "organ failure"]),
    ("death", ["death", "deaths", "died", "fatality", "fatalities", "mortality"]),
    ("cancer/carcinogenicity", ["cancer", "carcinogenicity", "carcinogenic", "cancers"]),
    ("vomiting", ["vomiting", "vomit"]),
    ("diarrhoea", ["diarrhoea", "diarrhea"]),
    ("food poisoning", ["food poisoning", "poisoning"]),
    ("gastrointestinal illness", ["gastroenteritis", "gastrointestinal distress", "digestive issues", "stomach cramps", "abdominal pain"]),
    ("allergic reaction", ["allergic reaction", "allergic reactions", "allergies", "allergy"]),
    ("liver damage", ["liver damage", "hepatotoxicity"]),
    ("cardiovascular disease", ["cardiovascular disease", "cardiovascular diseases", "heart disease", "heart attack", "heart attacks"]),
    ("neurological effects", ["neurological disorder", "neurological disorders", "neurological effects", "neurotoxicity"]),
    ("nausea", ["nausea"]),
    ("anuria", ["anuria"]),
    ("obesity", ["obesity"]),
]

COLORS = {
    "navy": "#163A5F",
    "blue": "#2B6A9F",
    "cyan": "#57A6B3",
    "green": "#3B7A57",
    "gold": "#D8A33D",
    "red": "#C75353",
    "purple": "#8064A2",
    "gray": "#7B8794",
    "light": "#DCE4EA",
    "ink": "#1E2933",
}


def clean(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def norm(value: object) -> str:
    text = clean(value).casefold().replace("-", " ").replace("/", " ")
    text = re.sub(r"[^\w\s]", " ", text, flags=re.UNICODE)
    return " ".join(text.split())


def local_predicate(value: object) -> str:
    return norm(clean(value).rsplit(":", 1)[-1]).replace(" ", "")


def finite_float(value: object) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def contains_alias(value: object, alias: object) -> bool:
    value_norm = f" {norm(value)} "
    alias_norm = f" {norm(alias)} "
    return alias_norm in value_norm


def canonical_from_aliases(value: object, groups: list[tuple[str, list[str]]]) -> str:
    raw = clean(value)
    for label, aliases in groups:
        if any(contains_alias(raw, alias) for alias in aliases):
            return label
    return norm(raw) or "(blank)"


def canonical_food(value: object) -> str:
    return canonical_from_aliases(value, FOOD_ALIASES)


def canonical_adulterant(value: object) -> str:
    return canonical_from_aliases(value, ADULTERANT_ALIASES)


def canonical_effect(value: object) -> str:
    return canonical_from_aliases(value, EFFECT_ALIASES)


def food_pool(source_id: object) -> str:
    value = clean(source_id).casefold()
    if value.startswith("oil-"):
        return "Oil"
    if value.startswith("milk-"):
        return "Milk"
    if value.startswith("ghee-"):
        return "Ghee"
    return "Ghee" if value and not value.startswith("src_") else "FSSAI baseline"


def confidence_band(value: float | None) -> str:
    if value is None:
        return "missing"
    if value < 0.70:
        return "<0.70"
    if value < 0.80:
        return "0.70-0.79"
    if value < 0.90:
        return "0.80-0.89"
    if value < 0.95:
        return "0.90-0.94"
    return ">=0.95"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_id(*parts: object) -> str:
    value = "\x1f".join(clean(part) for part in parts)
    return hashlib.sha256(value.encode("utf-8")).hexdigest()[:20]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_jsonl(path: Path) -> list[dict]:
    records: list[dict] = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                records.append(json.loads(line))
    return records


def write_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def markdown_table(headers: list[str], rows: list[list[object]]) -> str:
    def cell(value: object) -> str:
        return clean(value).replace("|", "\\|")

    output = [
        "| " + " | ".join(cell(header) for header in headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    output.extend("| " + " | ".join(cell(value) for value in row) + " |" for row in rows)
    return "\n".join(output)


def load_inputs() -> tuple[list[dict[str, str]], list[dict[str, str]], dict[str, dict]]:
    baseline = read_csv(BASELINE_PATH)
    news_raw = read_csv(NEWS_PATH)
    if list(baseline[0]) != BASE_COLUMNS:
        raise ValueError(f"FSSAI baseline columns changed: {list(baseline[0])}")
    if list(news_raw[0]) != BASE_COLUMNS + ["status"]:
        raise ValueError(f"News columns changed: {list(news_raw[0])}")
    news = [{column: row.get(column, "") for column in BASE_COLUMNS} for row in news_raw]
    articles: dict[str, dict] = {}
    for path in ARTICLE_PATHS:
        for record in read_jsonl(path):
            article_id = clean(record.get("article_id"))
            if article_id:
                articles[article_id] = record
    return baseline, news, articles


def normalized_rows(
    baseline: list[dict[str, str]], news: list[dict[str, str]], articles: dict[str, dict]
) -> list[dict]:
    output: list[dict] = []
    for corpus, rows in [("FSSAI baseline", baseline), ("News", news)]:
        for index, row in enumerate(rows, start=1):
            article = articles.get(row["source_id"], {})
            confidence = finite_float(row.get("confidence"))
            output.append(
                {
                    "triplet_id": stable_id(
                        corpus,
                        row.get("source_id"),
                        row.get("snippet_id"),
                        row.get("subject"),
                        row.get("predicate"),
                        row.get("object"),
                        row.get("evidence_span"),
                        index,
                    ),
                    "corpus": corpus,
                    "food_pool": food_pool(row.get("source_id")) if corpus == "News" else "FSSAI baseline",
                    "source_id": row.get("source_id", ""),
                    "source_file_or_url": row.get("source_file", ""),
                    "article_title": clean(article.get("title")),
                    "publication_date": clean(article.get("date")),
                    "publisher": clean(article.get("source")),
                    "snippet_id": row.get("snippet_id", ""),
                    "chunk_index": row.get("chunk_index", ""),
                    "subject": row.get("subject", ""),
                    "subject_type": row.get("subject_type", ""),
                    "predicate": row.get("predicate", ""),
                    "object": row.get("object", ""),
                    "object_type": row.get("object_type", ""),
                    "confidence": confidence,
                    "confidence_band": confidence_band(confidence),
                    "evidence_span": row.get("evidence_span", ""),
                    "normalized_subject": norm(row.get("subject")),
                    "normalized_predicate": local_predicate(row.get("predicate")),
                    "normalized_object": norm(row.get("object")),
                }
            )
    return output


def food_adulterant_edges(rows: list[dict[str, str]], corpus: str) -> list[dict]:
    output: list[dict] = []
    for row in rows:
        predicate = local_predicate(row.get("predicate"))
        if predicate not in DIRECT_FOOD_RELATIONS:
            continue
        if predicate == "hasadulterant":
            food_raw, adulterant_raw = row.get("subject", ""), row.get("object", "")
        else:
            food_raw, adulterant_raw = row.get("object", ""), row.get("subject", "")
        confidence = finite_float(row.get("confidence"))
        adulterant = canonical_adulterant(adulterant_raw)
        output.append(
            {
                "edge_id": stable_id(corpus, row.get("source_id"), row.get("snippet_id"), food_raw, adulterant_raw, predicate),
                "corpus": corpus,
                "food_pool": food_pool(row.get("source_id")) if corpus == "News" else "FSSAI baseline",
                "source_id": row.get("source_id", ""),
                "snippet_id": row.get("snippet_id", ""),
                "source_file_or_url": row.get("source_file", ""),
                "food": canonical_food(food_raw),
                "food_surface": clean(food_raw),
                "adulterant": adulterant,
                "adulterant_surface": clean(adulterant_raw),
                "relation": predicate,
                "confidence": confidence,
                "evidence_span": row.get("evidence_span", ""),
                "specific_candidate": norm(adulterant) not in GENERIC_ADULTERANTS,
            }
        )
    return output


def health_effect_edges(rows: list[dict[str, str]], corpus: str, adulterant_universe: set[str]) -> list[dict]:
    output: list[dict] = []
    for row in rows:
        predicate = local_predicate(row.get("predicate"))
        if predicate not in EFFECT_RELATIONS:
            continue
        subject_type = norm(row.get("subject_type"))
        object_type = norm(row.get("object_type"))
        subject = clean(row.get("subject"))
        obj = clean(row.get("object"))
        if not any(token in object_type for token in ["healtheffect", "acuteeffect", "chroniceffect"]):
            if predicate != "causeseffect":
                continue
        candidate = canonical_adulterant(subject)
        is_adulterant_type = any(
            token in subject_type
            for token in ["adulterant", "ingredient", "chemical", "foodadditive"]
        ) and "food" not in subject_type
        is_known_candidate = candidate in adulterant_universe and canonical_food(subject) == norm(subject)
        source_role = "adulterant" if is_adulterant_type or is_known_candidate else "other"
        confidence = finite_float(row.get("confidence"))
        output.append(
            {
                "edge_id": stable_id(corpus, row.get("source_id"), row.get("snippet_id"), subject, predicate, obj),
                "corpus": corpus,
                "food_pool": food_pool(row.get("source_id")) if corpus == "News" else "FSSAI baseline",
                "source_id": row.get("source_id", ""),
                "snippet_id": row.get("snippet_id", ""),
                "source_file_or_url": row.get("source_file", ""),
                "cause_surface": subject,
                "cause": candidate if source_role == "adulterant" else norm(subject),
                "cause_type": row.get("subject_type", ""),
                "cause_role": source_role,
                "effect": canonical_effect(obj),
                "effect_surface": obj,
                "effect_type": row.get("object_type", ""),
                "relation": predicate,
                "confidence": confidence,
                "evidence_span": row.get("evidence_span", ""),
            }
        )
    return output


def aggregate_food_adulterants(edges: list[dict]) -> list[dict]:
    grouped: defaultdict[tuple[str, str, str], list[dict]] = defaultdict(list)
    for edge in edges:
        grouped[(edge["corpus"], edge["food"], edge["adulterant"])].append(edge)
    output: list[dict] = []
    for (corpus, food, adulterant), group in grouped.items():
        confidences = [row["confidence"] for row in group if row["confidence"] is not None]
        source_confidence: defaultdict[str, list[float]] = defaultdict(list)
        for row in group:
            if row["confidence"] is not None:
                source_confidence[row["source_id"]].append(row["confidence"])
        source_max = [max(values) for values in source_confidence.values() if values]
        example = max(group, key=lambda row: row["confidence"] or 0)
        output.append(
            {
                "corpus": corpus,
                "food": food,
                "adulterant_or_substitute": adulterant,
                "specific_candidate": all(row["specific_candidate"] for row in group),
                "triplet_rows": len(group),
                "distinct_sources": len({row["source_id"] for row in group}),
                "mean_confidence": round(mean(confidences), 4) if confidences else "",
                "median_confidence": round(median(confidences), 4) if confidences else "",
                "minimum_confidence": round(min(confidences), 4) if confidences else "",
                "maximum_confidence": round(max(confidences), 4) if confidences else "",
                "confidence_weighted_source_support": round(sum(source_max), 4),
                "rows_confidence_ge_0_80": sum((row["confidence"] or 0) >= 0.80 for row in group),
                "rows_confidence_ge_0_90": sum((row["confidence"] or 0) >= 0.90 for row in group),
                "surface_forms": "; ".join(
                    f"{value} ({count})" for value, count in Counter(row["adulterant_surface"] for row in group).most_common(8)
                ),
                "example_source_id": example["source_id"],
                "example_url": example["source_file_or_url"],
                "example_evidence": example["evidence_span"],
            }
        )
    return sorted(
        output,
        key=lambda row: (
            row["corpus"],
            -int(row["distinct_sources"]),
            -float(row["mean_confidence"] or 0),
            row["food"],
            row["adulterant_or_substitute"],
        ),
    )


def relation_summary(rows: list[dict[str, str]], corpus: str) -> list[dict]:
    grouped: defaultdict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        grouped[local_predicate(row.get("predicate"))].append(row)
    output = []
    for predicate, group in grouped.items():
        values = [finite_float(row.get("confidence")) for row in group]
        values = [value for value in values if value is not None]
        output.append(
            {
                "corpus": corpus,
                "predicate": predicate,
                "triplet_rows": len(group),
                "distinct_sources": len({row.get("source_id", "") for row in group}),
                "mean_confidence": round(mean(values), 4) if values else "",
            }
        )
    return sorted(output, key=lambda row: (row["corpus"], -row["triplet_rows"]))


def build_gap_table(
    news_agg: list[dict],
    baseline_edges: list[dict],
    baseline: list[dict[str, str]],
    news: list[dict[str, str]],
    health_edges: list[dict],
) -> list[dict]:
    baseline_pairs = Counter((row["food"], row["adulterant"]) for row in baseline_edges)
    baseline_candidate_mentions: Counter[str] = Counter()
    for row in baseline:
        for field in ["subject", "object"]:
            baseline_candidate_mentions[canonical_adulterant(row.get(field, ""))] += 1

    baseline_methods_by_candidate: defaultdict[str, list[dict[str, str]]] = defaultdict(list)
    news_methods_by_candidate: defaultdict[str, list[dict[str, str]]] = defaultdict(list)
    for rows, index in [
        (baseline, baseline_methods_by_candidate),
        (news, news_methods_by_candidate),
    ]:
        for method_row in rows:
            if local_predicate(method_row.get("predicate")) not in METHOD_RELATIONS:
                continue
            candidates = {
                canonical_adulterant(method_row.get("subject")),
                canonical_adulterant(method_row.get("object")),
            }
            for candidate in candidates:
                index[candidate].append(method_row)

    health_by_candidate: defaultdict[str, list[dict]] = defaultdict(list)
    for row in health_edges:
        if row["cause_role"] == "adulterant":
            health_by_candidate[row["cause"]].append(row)

    news_specific = [
        row for row in news_agg if row["corpus"] == "News" and row["specific_candidate"]
    ]
    max_articles = max((int(row["distinct_sources"]) for row in news_specific), default=1)
    output = []
    for row in news_specific:
        food = row["food"]
        candidate = row["adulterant_or_substitute"]
        pair_rows = baseline_pairs[(food, candidate)]
        term_rows = baseline_candidate_mentions[candidate]
        baseline_methods = baseline_methods_by_candidate[candidate]
        news_methods = news_methods_by_candidate[candidate]
        health = health_by_candidate[candidate]
        health_articles = len({item["source_id"] for item in health if item["corpus"] == "News"})
        recurrence = math.log1p(int(row["distinct_sources"])) / math.log1p(max_articles)
        confidence = float(row["mean_confidence"] or 0)
        health_signal = min(1.0, health_articles / 5)
        priority = 100 * (
            0.35 * confidence
            + 0.25 * recurrence
            + 0.20 * health_signal
            + 0.10 * int(pair_rows == 0)
            + 0.10 * int(not baseline_methods)
        )
        if pair_rows:
            gap_class = "food-adulterant pair represented in FSSAI baseline"
        elif term_rows:
            gap_class = "candidate represented, but food-adulterant relation absent"
        else:
            gap_class = "candidate not represented in FSSAI baseline"
        if baseline_methods:
            test_class = "candidate-specific method relation represented"
        elif news_methods:
            test_class = "method signal occurs in news only"
        else:
            test_class = "no candidate-specific method relation represented"
        effects = Counter(item["effect"] for item in health)
        baseline_method_examples = Counter(
            f"{clean(item.get('subject'))} --{local_predicate(item.get('predicate'))}--> {clean(item.get('object'))}"
            for item in baseline_methods
        )
        news_method_examples = Counter(
            f"{clean(item.get('subject'))} --{local_predicate(item.get('predicate'))}--> {clean(item.get('object'))}"
            for item in news_methods
        )
        output.append(
            {
                "food": food,
                "adulterant_or_substitute": candidate,
                "news_triplet_rows": row["triplet_rows"],
                "news_distinct_articles": row["distinct_sources"],
                "mean_model_confidence": row["mean_confidence"],
                "confidence_weighted_article_support": row["confidence_weighted_source_support"],
                "fssai_baseline_direct_pair_rows": pair_rows,
                "fssai_baseline_candidate_mentions": term_rows,
                "fssai_baseline_candidate_method_rows": len(baseline_methods),
                "news_candidate_method_rows": len(news_methods),
                "fssai_baseline_method_examples": "; ".join(name for name, _ in baseline_method_examples.most_common(5)),
                "news_method_examples": "; ".join(name for name, _ in news_method_examples.most_common(5)),
                "news_health_relation_articles": health_articles,
                "locally_linked_effects": "; ".join(f"{name} ({count})" for name, count in effects.most_common(8)),
                "baseline_gap_class": gap_class,
                "test_coverage_class": test_class,
                "priority_score_0_100": round(priority, 2),
                "priority_interpretation": "heuristic discovery rank; not a prevalence or risk estimate",
                "example_url": row["example_url"],
                "example_evidence": row["example_evidence"],
            }
        )
    return sorted(output, key=lambda row: (-row["priority_score_0_100"], -row["news_distinct_articles"]))


def build_causal_tables(
    all_food_edges: list[dict], all_health_edges: list[dict], articles: dict[str, dict]
) -> tuple[list[dict], list[dict], list[dict], list[dict]]:
    food_group: defaultdict[tuple[str, str, str], list[dict]] = defaultdict(list)
    health_group: defaultdict[tuple[str, str, str], list[dict]] = defaultdict(list)
    for row in all_food_edges:
        food_group[(row["corpus"], row["food"], row["adulterant"])].append(row)
    for row in all_health_edges:
        if row["cause_role"] == "adulterant":
            health_group[(row["corpus"], row["cause"], row["effect"])].append(row)

    direct_same_article: defaultdict[tuple[str, str, str], set[str]] = defaultdict(set)
    direct_same_snippet: defaultdict[tuple[str, str, str], set[str]] = defaultdict(set)
    context_same_article: defaultdict[tuple[str, str, str], set[str]] = defaultdict(set)
    context_same_snippet: defaultdict[tuple[str, str, str], set[str]] = defaultdict(set)
    chain_article_confidence: defaultdict[tuple[str, str, str, str], list[float]] = defaultdict(list)
    direct_examples: defaultdict[tuple[str, str, str], list[tuple[dict, dict]]] = defaultdict(list)

    fa_by_source: defaultdict[tuple[str, str], list[dict]] = defaultdict(list)
    he_by_source: defaultdict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in all_food_edges:
        fa_by_source[(row["corpus"], row["source_id"])].append(row)
    for row in all_health_edges:
        he_by_source[(row["corpus"], row["source_id"])].append(row)

    for source_key, food_rows in fa_by_source.items():
        effect_rows = he_by_source.get(source_key, [])
        corpus, source_id = source_key
        for fa in food_rows:
            for he in effect_rows:
                key = (fa["food"], fa["adulterant"], he["effect"])
                values = [value for value in [fa["confidence"], he["confidence"]] if value is not None]
                chain_conf = min(values) if values else 0.0
                if he["cause_role"] == "adulterant" and he["cause"] == fa["adulterant"]:
                    direct_same_article[key].add(source_id)
                    direct_examples[key].append((fa, he))
                    if fa["snippet_id"] == he["snippet_id"]:
                        direct_same_snippet[key].add(source_id)
                    chain_article_confidence[(corpus, *key)].append(chain_conf)
                else:
                    context_same_article[key].add(source_id)
                    if fa["snippet_id"] == he["snippet_id"]:
                        context_same_snippet[key].add(source_id)

    fa_any: defaultdict[tuple[str, str], list[dict]] = defaultdict(list)
    ae_any: defaultdict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in all_food_edges:
        fa_any[(row["food"], row["adulterant"])].append(row)
    for row in all_health_edges:
        if row["cause_role"] == "adulterant":
            ae_any[(row["cause"], row["effect"])].append(row)

    chains: list[dict] = []
    effects_by_adulterant: defaultdict[str, list[tuple[str, list[dict]]]] = defaultdict(list)
    for (cause, effect), ae_rows in ae_any.items():
        effects_by_adulterant[cause].append((effect, ae_rows))
    for (food, adulterant), fa_rows in fa_any.items():
        for effect, ae_rows in effects_by_adulterant.get(adulterant, []):
            key = (food, adulterant, effect)
            fa_conf = [row["confidence"] for row in fa_rows if row["confidence"] is not None]
            ae_conf = [row["confidence"] for row in ae_rows if row["confidence"] is not None]
            fa_news_ids = {row["source_id"] for row in fa_rows if row["corpus"] == "News"}
            ae_news_ids = {row["source_id"] for row in ae_rows if row["corpus"] == "News"}
            same_direct = direct_same_article[key]
            examples = direct_examples[key]
            example_fa, example_he = max(
                examples,
                key=lambda pair: min(pair[0]["confidence"] or 0, pair[1]["confidence"] or 0),
            ) if examples else (None, None)
            if same_direct:
                chain_class = "directly joined in the same article"
            elif fa_news_ids & ae_news_ids:
                chain_class = "news article-level join"
            else:
                chain_class = "repository-inferred from separate sources"
            chains.append(
                {
                    "food": food,
                    "adulterant_or_substitute": adulterant,
                    "effect": effect,
                    "chain_class": chain_class,
                    "food_adulterant_news_articles": len(fa_news_ids),
                    "adulterant_effect_news_articles": len(ae_news_ids),
                    "same_article_direct_chain_articles": len(same_direct),
                    "same_snippet_direct_chain_articles": len(direct_same_snippet[key]),
                    "same_article_context_articles": len(context_same_article[key]),
                    "fssai_food_adulterant_rows": sum(row["corpus"] == "FSSAI baseline" for row in fa_rows),
                    "fssai_adulterant_effect_rows": sum(row["corpus"] == "FSSAI baseline" for row in ae_rows),
                    "mean_food_adulterant_confidence": round(mean(fa_conf), 4) if fa_conf else "",
                    "mean_adulterant_effect_confidence": round(mean(ae_conf), 4) if ae_conf else "",
                    "conservative_chain_confidence": round(min(mean(fa_conf or [0]), mean(ae_conf or [0])), 4),
                    "chain_support_score": round(
                        min(mean(fa_conf or [0]), mean(ae_conf or [0]))
                        * math.log1p(len(fa_news_ids) + len(ae_news_ids) + 2 * len(same_direct)),
                        4,
                    ),
                    "example_article_id": example_fa["source_id"] if example_fa else "",
                    "example_article_url": example_fa["source_file_or_url"] if example_fa else "",
                    "example_food_adulterant_evidence": example_fa["evidence_span"] if example_fa else "",
                    "example_adulterant_effect_evidence": example_he["evidence_span"] if example_he else "",
                    "interpretation": "reported or repository-inferred association; not epidemiological proof of causation",
                }
            )
    chains.sort(key=lambda row: (-row["same_article_direct_chain_articles"], -row["chain_support_score"]))

    contextual: list[dict] = []
    for key, ids in context_same_article.items():
        if not ids:
            continue
        food, adulterant, effect = key
        contextual.append(
            {
                "food": food,
                "adulterant_or_substitute": adulterant,
                "effect": effect,
                "same_article_articles": len(ids),
                "same_snippet_articles": len(context_same_snippet[key]),
                "article_ids": "; ".join(sorted(ids)[:25]),
                "inference_level": "same-article contextual hypothesis; causal edge is not explicit",
            }
        )
    contextual.sort(key=lambda row: (-row["same_snippet_articles"], -row["same_article_articles"]))

    news_food_edges = [row for row in all_food_edges if row["corpus"] == "News"]
    news_health = [row for row in all_health_edges if row["corpus"] == "News"]
    fa_by_article: defaultdict[str, list[dict]] = defaultdict(list)
    effects_by_article: defaultdict[str, set[str]] = defaultdict(set)
    for row in news_food_edges:
        fa_by_article[row["source_id"]].append(row)
    for row in news_health:
        effects_by_article[row["source_id"]].add(row["effect"])

    denominator: Counter[tuple[str, str]] = Counter()
    numerator: Counter[tuple[str, str, str]] = Counter()
    numerator_conf: defaultdict[tuple[str, str, str], list[float]] = defaultdict(list)
    for article_id, fa_rows in fa_by_article.items():
        effects = effects_by_article.get(article_id, set())
        food_candidates: defaultdict[str, list[dict]] = defaultdict(list)
        for row in fa_rows:
            food_candidates[row["food"]].append(row)
        for food, candidate_rows in food_candidates.items():
            for effect in effects:
                denominator[(food, effect)] += 1
                for candidate in {row["adulterant"] for row in candidate_rows}:
                    numerator[(food, effect, candidate)] += 1
                    confidence_values = [
                        row["confidence"] or 0
                        for row in candidate_rows
                        if row["adulterant"] == candidate
                    ]
                    numerator_conf[(food, effect, candidate)].append(max(confidence_values or [0]))

    reverse: list[dict] = []
    for (food, effect, candidate), count in numerator.items():
        total = denominator[(food, effect)]
        reverse.append(
            {
                "food": food,
                "observed_effect": effect,
                "candidate_adulterant": candidate,
                "articles_with_food_effect_context": total,
                "articles_also_naming_candidate": count,
                "observed_article_share_pct": round(100 * count / total, 2) if total else 0,
                "laplace_smoothed_presence_pct": round(100 * (count + 1) / (total + 2), 2),
                "mean_candidate_confidence": round(mean(numerator_conf[(food, effect, candidate)]), 4),
                "interpretation": "hypothesis-ranking association, not a calibrated causal probability",
            }
        )
    reverse.sort(key=lambda row: (-row["articles_also_naming_candidate"], -row["observed_article_share_pct"]))

    specificity_ids: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    for row in news_food_edges:
        specificity_ids[(row["adulterant"], row["food"])].add(row["source_id"])
    specificity: list[dict] = []
    candidate_totals = Counter()
    for (candidate, food), article_ids in specificity_ids.items():
        candidate_totals[candidate] += len(article_ids)
    for (candidate, food), article_ids in specificity_ids.items():
        count = len(article_ids)
        specificity.append(
            {
                "adulterant_or_substitute": candidate,
                "food": food,
                "distinct_articles": count,
                "share_of_candidate_food_links_pct": round(100 * count / candidate_totals[candidate], 2),
                "candidate_total_food_link_articles": candidate_totals[candidate],
            }
        )
    specificity.sort(key=lambda row: (-row["candidate_total_food_link_articles"], -row["distinct_articles"]))
    return chains, contextual, reverse, specificity


def build_temporal_table(edges: list[dict], articles: dict[str, dict]) -> list[dict]:
    grouped: defaultdict[tuple[int, str, str], list[dict]] = defaultdict(list)
    for row in edges:
        if row["corpus"] != "News":
            continue
        date_text = clean(articles.get(row["source_id"], {}).get("date"))
        match = re.match(r"^(\d{4})", date_text)
        if not match:
            continue
        grouped[(int(match.group(1)), row["food"], row["adulterant"])].append(row)
    output = []
    for (year, food, candidate), rows in grouped.items():
        ids = {row["source_id"] for row in rows}
        values = [row["confidence"] for row in rows if row["confidence"] is not None]
        output.append(
            {
                "publication_year": year,
                "food": food,
                "adulterant_or_substitute": candidate,
                "triplet_rows": len(rows),
                "distinct_articles": len(ids),
                "mean_confidence": round(mean(values), 4) if values else "",
            }
        )
    return sorted(output, key=lambda row: (row["publication_year"], row["food"], -row["distinct_articles"]))


def core_profiles(
    baseline: list[dict[str, str]], news: list[dict[str, str]], articles: dict[str, dict]
) -> tuple[list[dict], list[dict], dict]:
    profiles = []
    for corpus, rows in [("FSSAI baseline", baseline), ("News", news)]:
        confidence = [finite_float(row.get("confidence")) for row in rows]
        confidence = [value for value in confidence if value is not None]
        triples = {
            (norm(row.get("subject")), local_predicate(row.get("predicate")), norm(row.get("object")))
            for row in rows
        }
        entities = {norm(row.get("subject")) for row in rows} | {norm(row.get("object")) for row in rows}
        profiles.append(
            {
                "corpus": corpus,
                "triplet_rows": len(rows),
                "unique_normalized_triplets": len(triples),
                "unique_entities": len(entities),
                "distinct_sources": len({row.get("source_id", "") for row in rows}),
                "mean_confidence": round(mean(confidence), 4),
                "median_confidence": round(median(confidence), 4),
                "minimum_confidence": min(confidence),
                "maximum_confidence": max(confidence),
            }
        )
    baseline_triples = {
        (norm(row["subject"]), local_predicate(row["predicate"]), norm(row["object"])) for row in baseline
    }
    news_triples = {
        (norm(row["subject"]), local_predicate(row["predicate"]), norm(row["object"])) for row in news
    }
    baseline_entities = {norm(row["subject"]) for row in baseline} | {norm(row["object"]) for row in baseline}
    news_entities = {norm(row["subject"]) for row in news} | {norm(row["object"]) for row in news}
    overlap = {
        "exact_normalized_triplet_overlap": len(baseline_triples & news_triples),
        "entity_overlap": len(baseline_entities & news_entities),
        "news_only_normalized_triplets": len(news_triples - baseline_triples),
        "baseline_only_normalized_triplets": len(baseline_triples - news_triples),
        "news_only_entities": len(news_entities - baseline_entities),
        "baseline_only_entities": len(baseline_entities - news_entities),
    }
    food_rows = []
    for food in ["Oil", "Ghee", "Milk"]:
        rows = [row for row in news if food_pool(row["source_id"]) == food]
        values = [finite_float(row["confidence"]) for row in rows]
        values = [value for value in values if value is not None]
        ids = {row["source_id"] for row in rows}
        input_ids = {article_id for article_id in articles if food_pool(article_id) == food}
        food_rows.append(
            {
                "food_pool": food,
                "input_articles": len(input_ids),
                "articles_with_triplets": len(ids),
                "triplet_rows": len(rows),
                "mean_confidence": round(mean(values), 4),
                "median_confidence": round(median(values), 4),
                "rows_confidence_ge_0_90": sum(value >= 0.90 for value in values),
            }
        )
    return profiles, food_rows, overlap


def aggregate_health_edges(edges: list[dict]) -> list[dict]:
    grouped: defaultdict[tuple[str, str, str, str], list[dict]] = defaultdict(list)
    for row in edges:
        grouped[(row["corpus"], row["food_pool"], row["cause"], row["effect"])].append(row)
    output = []
    for (corpus, pool_name, cause, effect), rows in grouped.items():
        values = [row["confidence"] for row in rows if row["confidence"] is not None]
        example = max(rows, key=lambda row: row["confidence"] or 0)
        output.append(
            {
                "corpus": corpus,
                "food_pool": pool_name,
                "cause": cause,
                "cause_role": Counter(row["cause_role"] for row in rows).most_common(1)[0][0],
                "effect": effect,
                "triplet_rows": len(rows),
                "distinct_sources": len({row["source_id"] for row in rows}),
                "mean_confidence": round(mean(values), 4) if values else "",
                "example_url": example["source_file_or_url"],
                "example_evidence": example["evidence_span"],
            }
        )
    return sorted(output, key=lambda row: (row["corpus"], -row["distinct_sources"], -row["triplet_rows"]))


def aggregate_selected_relations(
    rows: list[dict[str, str]], predicates: set[str], corpus: str
) -> list[dict]:
    grouped: defaultdict[tuple[str, str, str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        pred = local_predicate(row.get("predicate"))
        if pred not in predicates:
            continue
        grouped[(food_pool(row.get("source_id")) if corpus == "News" else "FSSAI baseline", clean(row.get("subject")), pred, clean(row.get("object")))].append(row)
    output = []
    for (pool_name, subject, pred, obj), group in grouped.items():
        values = [finite_float(row.get("confidence")) for row in group]
        values = [value for value in values if value is not None]
        example = max(group, key=lambda row: finite_float(row.get("confidence")) or 0)
        output.append(
            {
                "corpus": corpus,
                "food_pool": pool_name,
                "subject": subject,
                "predicate": pred,
                "object": obj,
                "triplet_rows": len(group),
                "distinct_sources": len({row.get("source_id", "") for row in group}),
                "mean_confidence": round(mean(values), 4) if values else "",
                "example_url": example.get("source_file", ""),
                "example_evidence": example.get("evidence_span", ""),
            }
        )
    return sorted(output, key=lambda row: (row["corpus"], -row["distinct_sources"], -row["triplet_rows"]))


def confidence_sensitivity(news: list[dict[str, str]]) -> list[dict]:
    output = []
    for pool_name in ["Oil", "Ghee", "Milk", "All news"]:
        pool_rows = news if pool_name == "All news" else [row for row in news if food_pool(row["source_id"]) == pool_name]
        for threshold in [0.0, 0.70, 0.80, 0.90, 0.95]:
            selected = [row for row in pool_rows if (finite_float(row.get("confidence")) or 0) >= threshold]
            output.append(
                {
                    "food_pool": pool_name,
                    "minimum_confidence": threshold,
                    "triplet_rows": len(selected),
                    "distinct_articles": len({row["source_id"] for row in selected}),
                    "share_of_pool_rows_pct": round(100 * len(selected) / len(pool_rows), 2) if pool_rows else 0,
                }
            )
    return output


def publisher_summary(news: list[dict[str, str]], articles: dict[str, dict]) -> list[dict]:
    ids_by_publisher: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    rows_by_publisher: Counter[tuple[str, str]] = Counter()
    for row in news:
        article_id = row["source_id"]
        pool_name = food_pool(article_id)
        publisher = clean(articles.get(article_id, {}).get("source")) or "(missing publisher)"
        ids_by_publisher[(pool_name, publisher)].add(article_id)
        rows_by_publisher[(pool_name, publisher)] += 1
    output = [
        {
            "food_pool": pool_name,
            "publisher": publisher,
            "distinct_articles": len(article_ids),
            "triplet_rows": rows_by_publisher[(pool_name, publisher)],
        }
        for (pool_name, publisher), article_ids in ids_by_publisher.items()
    ]
    return sorted(output, key=lambda row: (row["food_pool"], -row["distinct_articles"], row["publisher"]))


def build_test_priorities(gaps: list[dict]) -> list[dict]:
    output = []
    for row in gaps:
        if row["fssai_baseline_candidate_method_rows"]:
            method_basis = "Candidate-specific method relation exists in the FSSAI baseline."
            next_step = "Retrieve the represented procedure and validate it for the exact food matrix, analyte range and enforcement purpose."
        elif row["news_candidate_method_rows"]:
            method_basis = "Candidate-specific method signal occurs in news triplets but not in the FSSAI baseline graph."
            next_step = "Treat the news method as a retrieval lead; verify the full protocol locally before laboratory or enforcement use."
        else:
            method_basis = "No candidate-specific method relation is represented in either supplied triplet corpus."
            next_step = "Prioritize method discovery and matrix-specific validation; this dataset cannot responsibly name a test."
        health_basis = row["locally_linked_effects"] or "No candidate-specific health effect is linked in the local graph."
        output.append(
            {
                "priority_rank": 0,
                "food": row["food"],
                "candidate": row["adulterant_or_substitute"],
                "priority_score_0_100": row["priority_score_0_100"],
                "news_distinct_articles": row["news_distinct_articles"],
                "mean_model_confidence": row["mean_model_confidence"],
                "baseline_gap_class": row["baseline_gap_class"],
                "method_basis_from_local_graph": method_basis,
                "fssai_baseline_method_examples": row["fssai_baseline_method_examples"],
                "news_method_examples": row["news_method_examples"],
                "local_health_basis": health_basis,
                "recommended_action": next_step,
                "minimum_validation_record": "food matrix; analyte identity; sampling plan; blanks; spikes; recovery; precision; detection and quantification limits; confirmatory rule; chain of custody",
                "interpretation_boundary": "testing recommendation derived from local graph coverage; not a validated analytical protocol",
            }
        )
    output.sort(key=lambda row: (-row["priority_score_0_100"], -row["news_distinct_articles"]))
    for index, row in enumerate(output, start=1):
        row["priority_rank"] = index
    return output


def build_all_tables() -> dict[str, object]:
    baseline, news, articles = load_inputs()
    normalized = normalized_rows(baseline, news, articles)
    baseline_food_edges = food_adulterant_edges(baseline, "FSSAI baseline")
    news_food_edges = food_adulterant_edges(news, "News")
    all_food_edges = baseline_food_edges + news_food_edges
    adulterant_universe = {row["adulterant"] for row in all_food_edges}
    baseline_health = health_effect_edges(baseline, "FSSAI baseline", adulterant_universe)
    news_health = health_effect_edges(news, "News", adulterant_universe)
    all_health = baseline_health + news_health
    food_agg = aggregate_food_adulterants(all_food_edges)
    gaps = build_gap_table(food_agg, baseline_food_edges, baseline, news, all_health)
    chains, contextual, reverse, specificity = build_causal_tables(all_food_edges, all_health, articles)
    profiles, food_profiles, overlap = core_profiles(baseline, news, articles)
    temporal = build_temporal_table(news_food_edges, articles)
    predicates = relation_summary(baseline, "FSSAI baseline") + relation_summary(news, "News")
    health_summary = aggregate_health_edges(all_health)
    methods = aggregate_selected_relations(baseline, METHOD_RELATIONS, "FSSAI baseline") + aggregate_selected_relations(news, METHOD_RELATIONS, "News")
    actions = aggregate_selected_relations(baseline, ACTION_RELATIONS, "FSSAI baseline") + aggregate_selected_relations(news, ACTION_RELATIONS, "News")
    sensitivity = confidence_sensitivity(news)
    publishers = publisher_summary(news, articles)
    test_priorities = build_test_priorities(gaps)

    return {
        "baseline": baseline,
        "news": news,
        "articles": articles,
        "normalized": normalized,
        "baseline_food_edges": baseline_food_edges,
        "news_food_edges": news_food_edges,
        "all_food_edges": all_food_edges,
        "baseline_health_edges": baseline_health,
        "news_health_edges": news_health,
        "all_health_edges": all_health,
        "food_adulterants": food_agg,
        "gaps": gaps,
        "causal_chains": chains,
        "contextual_chains": contextual,
        "reverse_hypotheses": reverse,
        "food_specificity": specificity,
        "profiles": profiles,
        "food_profiles": food_profiles,
        "overlap": overlap,
        "temporal": temporal,
        "predicates": predicates,
        "health_summary": health_summary,
        "methods": methods,
        "actions": actions,
        "confidence_sensitivity": sensitivity,
        "publishers": publishers,
        "test_priorities": test_priorities,
    }


def set_plot_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9.5,
            "axes.titlesize": 14,
            "axes.titleweight": "bold",
            "axes.labelsize": 10,
            "axes.edgecolor": COLORS["gray"],
            "axes.grid": True,
            "axes.axisbelow": True,
            "grid.color": "#D8E0E6",
            "grid.linewidth": 0.7,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "legend.frameon": False,
            "savefig.bbox": "tight",
        }
    )


class FigureWriter:
    def __init__(self) -> None:
        FIGURES.mkdir(parents=True, exist_ok=True)
        self.pdf = PdfPages(OUT / f"FSSAI_Triplet_Intelligence_Figures_{DATE_TAG}.pdf")
        self.manifest: list[dict] = []
        self.number = 0

    def save(self, fig, slug: str, title: str, caption: str, data: pd.DataFrame) -> None:
        self.number += 1
        stem = f"figure_{self.number:02d}_{slug}"
        png = FIGURES / f"{stem}.png"
        data_path = TABLES / f"{stem}_data.csv"
        data.to_csv(data_path, index=False, encoding="utf-8-sig")
        fig.savefig(png, dpi=190, facecolor="white")
        self.pdf.savefig(fig, bbox_inches="tight", facecolor="white")
        plt.close(fig)
        self.manifest.append(
            {
                "number": self.number,
                "slug": slug,
                "title": title,
                "caption": caption,
                "png": str(png.relative_to(OUT)),
                "data_csv": str(data_path.relative_to(OUT)),
            }
        )

    def close(self) -> list[dict]:
        self.pdf.close()
        (OUT / "figure_manifest.json").write_text(
            json.dumps(self.manifest, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        return self.manifest


def horizontal_bar(
    writer: FigureWriter,
    data: pd.DataFrame,
    label: str,
    value: str,
    slug: str,
    title: str,
    caption: str,
    color: str = COLORS["blue"],
    value_format: str = "{:.0f}",
) -> None:
    frame = data.copy().sort_values(value, ascending=True)
    fig, ax = plt.subplots(figsize=(10.6, 6.3))
    bars = ax.barh(frame[label], frame[value], color=color)
    ax.set_title(title, loc="left", pad=14)
    ax.set_xlabel(value.replace("_", " ").title())
    ax.grid(axis="y", visible=False)
    ax.spines[["top", "right", "left"]].set_visible(False)
    limit = max(frame[value].max(), 1)
    for bar, item in zip(bars, frame[value]):
        ax.text(item + limit * 0.012, bar.get_y() + bar.get_height() / 2, value_format.format(item), va="center", fontsize=8)
    ax.set_xlim(0, limit * 1.14)
    fig.tight_layout()
    writer.save(fig, slug, title, caption, frame)


def generate_figures(tables: dict[str, object]) -> list[dict]:
    set_plot_style()
    writer = FigureWriter()
    profiles = pd.DataFrame(tables["profiles"])
    food_profiles = pd.DataFrame(tables["food_profiles"])
    news = tables["news"]
    baseline = tables["baseline"]
    news_edges = tables["news_food_edges"]
    gaps = pd.DataFrame(tables["gaps"])
    chains = pd.DataFrame(tables["causal_chains"])
    reverse = pd.DataFrame(tables["reverse_hypotheses"])

    # 1. Corpus rows and unique triplets.
    plot = profiles[["corpus", "triplet_rows", "unique_normalized_triplets"]].copy()
    x = np.arange(len(plot))
    fig, ax = plt.subplots(figsize=(9.4, 5.8))
    ax.bar(x - 0.18, plot["triplet_rows"], 0.36, label="Triplet rows", color=COLORS["navy"])
    ax.bar(x + 0.18, plot["unique_normalized_triplets"], 0.36, label="Unique normalized triplets", color=COLORS["cyan"])
    ax.set_xticks(x, plot["corpus"])
    ax.set_ylabel("Count")
    ax.set_title("Comparable corpus scale", loc="left")
    ax.legend()
    ax.grid(axis="x", visible=False)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    writer.save(fig, "corpus_scale", "Comparable corpus scale", "All retained news triplets are included; no status-based partition is used.", plot)

    # 2. Entities and relations.
    plot = profiles[["corpus", "unique_entities", "unique_normalized_triplets"]].copy()
    fig, ax = plt.subplots(figsize=(9.4, 5.8))
    ax.scatter(plot["unique_entities"], plot["unique_normalized_triplets"], s=[260, 260], c=[COLORS["navy"], COLORS["green"]])
    for _, row in plot.iterrows():
        ax.annotate(row["corpus"], (row["unique_entities"], row["unique_normalized_triplets"]), xytext=(8, 8), textcoords="offset points")
    ax.set_xlabel("Unique normalized entities")
    ax.set_ylabel("Unique normalized triplets")
    ax.set_title("Knowledge-graph breadth", loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    writer.save(fig, "graph_breadth", "Knowledge-graph breadth", "The corpora are similar in row scale but encode different knowledge functions.", plot)

    # 3. Confidence distributions.
    bconf = np.array([finite_float(row["confidence"]) for row in baseline], dtype=float)
    nconf = np.array([finite_float(row["confidence"]) for row in news], dtype=float)
    bins = np.arange(0.45, 1.001, 0.025)
    fig, ax = plt.subplots(figsize=(10.2, 5.9))
    ax.hist(bconf, bins=bins, alpha=0.60, label="FSSAI baseline", color=COLORS["navy"], density=True)
    ax.hist(nconf, bins=bins, alpha=0.58, label="News", color=COLORS["green"], density=True)
    ax.axvline(np.median(nconf), color=COLORS["green"], linestyle="--", linewidth=1.5)
    ax.set_xlabel("Model confidence")
    ax.set_ylabel("Density")
    ax.set_title("Confidence distributions", loc="left")
    ax.legend()
    ax.spines[["top", "right"]].set_visible(False)
    data = pd.DataFrame({"corpus": ["FSSAI baseline", "News"], "rows": [len(bconf), len(nconf)], "mean": [bconf.mean(), nconf.mean()], "median": [np.median(bconf), np.median(nconf)]})
    fig.tight_layout()
    writer.save(fig, "confidence_distribution", "Confidence distributions", "Confidence is preserved as supplied and is the only model-quality signal used for news rankings.", data)

    # 4. News confidence by food pool.
    confidence_by_food = [[finite_float(row["confidence"]) for row in news if food_pool(row["source_id"]) == food] for food in ["Oil", "Ghee", "Milk"]]
    fig, ax = plt.subplots(figsize=(9.4, 5.8))
    bp = ax.boxplot(confidence_by_food, labels=["Oil", "Ghee", "Milk"], patch_artist=True, showfliers=False)
    for patch, color in zip(bp["boxes"], [COLORS["gold"], COLORS["green"], COLORS["blue"]]):
        patch.set_facecolor(color)
        patch.set_alpha(0.78)
    ax.set_ylabel("Model confidence")
    ax.set_ylim(0.55, 1.0)
    ax.set_title("News confidence by food pool", loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    writer.save(fig, "confidence_by_food", "News confidence by food pool", "The three pools have similar confidence distributions; counts and food-specific evidence drive substantive differences.", food_profiles)

    # 5. Rows and represented articles by pool.
    plot = food_profiles.copy()
    fig, ax = plt.subplots(figsize=(9.6, 5.9))
    x = np.arange(len(plot))
    ax.bar(x - 0.18, plot["triplet_rows"], 0.36, color=COLORS["navy"], label="Triplet rows")
    ax2 = ax.twinx()
    ax2.bar(x + 0.18, plot["articles_with_triplets"], 0.36, color=COLORS["gold"], label="Articles with triplets")
    ax.set_xticks(x, plot["food_pool"])
    ax.set_ylabel("Triplet rows")
    ax2.set_ylabel("Articles")
    ax.set_title("News-triplet coverage by food pool", loc="left")
    handles = ax.containers + ax2.containers
    ax.legend([item for container in handles for item in [container]], ["Triplet rows", "Articles with triplets"], loc="upper right")
    ax.grid(axis="x", visible=False)
    ax2.grid(False)
    fig.tight_layout()
    writer.save(fig, "food_pool_coverage", "News-triplet coverage by food pool", "Article counts measure corpus representation, not incident prevalence.", plot)

    # 6. Predicate comparison.
    predicate_frame = pd.DataFrame(tables["predicates"])
    top_names = list(predicate_frame.groupby("predicate")["triplet_rows"].sum().nlargest(14).index)
    pivot = predicate_frame[predicate_frame["predicate"].isin(top_names)].pivot_table(index="predicate", columns="corpus", values="triplet_rows", aggfunc="sum", fill_value=0)
    pivot = pivot.loc[pivot.sum(axis=1).sort_values().index]
    fig, ax = plt.subplots(figsize=(10.6, 6.7))
    y = np.arange(len(pivot))
    ax.barh(y - 0.18, pivot.get("FSSAI baseline", 0), 0.36, label="FSSAI baseline", color=COLORS["navy"])
    ax.barh(y + 0.18, pivot.get("News", 0), 0.36, label="News", color=COLORS["green"])
    ax.set_yticks(y, pivot.index)
    ax.set_xscale("symlog", linthresh=10)
    ax.set_xlabel("Triplet rows (symmetric log scale)")
    ax.set_title("Dominant predicates reveal different corpus functions", loc="left")
    ax.legend()
    ax.grid(axis="y", visible=False)
    fig.tight_layout()
    writer.save(fig, "predicate_comparison", "Dominant predicates reveal different corpus functions", "FSSAI baseline emphasizes standards and definitions; news emphasizes incidents, adulterants, effects, places and actions.", pivot.reset_index())

    # 7. Exact overlap.
    overlap = tables["overlap"]
    plot = pd.DataFrame(
        [
            ["Exact shared triplets", overlap["exact_normalized_triplet_overlap"]],
            ["News-only triplets", overlap["news_only_normalized_triplets"]],
            ["FSSAI-baseline-only triplets", overlap["baseline_only_normalized_triplets"]],
            ["Shared entities", overlap["entity_overlap"]],
        ],
        columns=["category", "count"],
    )
    horizontal_bar(writer, plot, "category", "count", "overlap", "Normalized overlap is sparse", "Exact overlap is conservative; terminology and role differences can hide contextual overlap.", COLORS["purple"])

    def aggregate_candidate_articles(pool_filter: str | None = None) -> pd.DataFrame:
        grouped: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
        confidence: defaultdict[tuple[str, str], list[float]] = defaultdict(list)
        for row in news_edges:
            if pool_filter and row["food_pool"] != pool_filter:
                continue
            if not row["specific_candidate"]:
                continue
            key = (row["food"], row["adulterant"])
            grouped[key].add(row["source_id"])
            if row["confidence"] is not None:
                confidence[key].append(row["confidence"])
        return pd.DataFrame(
            [
                {"food": key[0], "candidate": key[1], "food_candidate": f"{key[0]} -> {key[1]}", "distinct_articles": len(ids), "mean_confidence": mean(confidence[key])}
                for key, ids in grouped.items()
            ]
        )

    # 8-11. Overall and food-pool rankings.
    overall = aggregate_candidate_articles().nlargest(18, "distinct_articles")
    horizontal_bar(writer, overall, "food_candidate", "distinct_articles", "top_food_adulterants", "Most reported food-adulterant relationships", "Ranks use unique articles and retain the explicitly named affected food.", COLORS["red"])
    for pool_name, color in [("Oil", COLORS["gold"]), ("Ghee", COLORS["green"]), ("Milk", COLORS["blue"])]:
        frame = aggregate_candidate_articles(pool_name).nlargest(15, "distinct_articles")
        horizontal_bar(writer, frame, "food_candidate", "distinct_articles", f"top_{pool_name.casefold()}", f"Most reported relationships in the {pool_name} article pool", "Food is taken from the triplet; the pool only identifies the source collection.", color)

    # 12. Food-adulterant heatmap.
    pair_frame = aggregate_candidate_articles()
    foods = list(pair_frame.groupby("food")["distinct_articles"].sum().nlargest(10).index)
    candidates = list(pair_frame.groupby("candidate")["distinct_articles"].sum().nlargest(14).index)
    heat = pair_frame[pair_frame["food"].isin(foods) & pair_frame["candidate"].isin(candidates)].pivot_table(index="candidate", columns="food", values="distinct_articles", aggfunc="sum", fill_value=0)
    fig, ax = plt.subplots(figsize=(11.5, 7.0))
    image = ax.imshow(heat.values, aspect="auto", cmap="YlOrRd")
    ax.set_xticks(np.arange(len(heat.columns)), heat.columns, rotation=35, ha="right")
    ax.set_yticks(np.arange(len(heat.index)), heat.index)
    ax.set_title("Food-specific adulterant article matrix", loc="left")
    fig.colorbar(image, ax=ax, label="Distinct articles")
    fig.tight_layout()
    writer.save(fig, "food_adulterant_heatmap", "Food-specific adulterant article matrix", "The same candidate is not assumed equally likely across foods.", heat.reset_index())

    # 13. Recurrence versus confidence.
    scatter = pair_frame[pair_frame["distinct_articles"] >= 2].copy()
    fig, ax = plt.subplots(figsize=(10.3, 6.2))
    colors = [COLORS["blue"] if food.startswith("Milk") else COLORS["green"] if food == "Ghee" else COLORS["gold"] for food in scatter["food"]]
    ax.scatter(scatter["distinct_articles"], scatter["mean_confidence"], c=colors, alpha=0.68, s=np.clip(scatter["distinct_articles"] * 4, 18, 240))
    for _, row in scatter.nlargest(12, "distinct_articles").iterrows():
        ax.annotate(row["food_candidate"], (row["distinct_articles"], row["mean_confidence"]), xytext=(4, 4), textcoords="offset points", fontsize=7)
    ax.set_xscale("log")
    ax.set_xlabel("Distinct articles (log scale)")
    ax.set_ylabel("Mean model confidence")
    ax.set_title("Recurrence and confidence are complementary signals", loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    writer.save(fig, "recurrence_confidence", "Recurrence and confidence are complementary signals", "Confidence ranks extraction certainty; recurrence measures corpus repetition and can include duplicate incident coverage.", scatter)

    # 14. Gap classes.
    gap_counts = gaps.groupby("baseline_gap_class", as_index=False).agg(food_candidate_pairs=("adulterant_or_substitute", "size"), news_articles=("news_distinct_articles", "sum"))
    horizontal_bar(writer, gap_counts, "baseline_gap_class", "food_candidate_pairs", "gap_classes", "FSSAI baseline representation classes", "Absence means absent from the supplied verified baseline, not necessarily from every current FSSAI document.", COLORS["purple"])

    # 15. Priority gap candidates.
    top_gap = gaps.nlargest(18, "priority_score_0_100").copy()
    top_gap["label"] = top_gap["food"] + " -> " + top_gap["adulterant_or_substitute"]
    top_gap = top_gap.sort_values("priority_score_0_100", ascending=True).reset_index(drop=True)
    y_positions = np.arange(len(top_gap))
    marker_sizes = np.clip(np.log1p(top_gap["news_distinct_articles"]) * 90, 70, 470)
    fig, ax = plt.subplots(figsize=(11.3, 7.4))
    scatter_obj = ax.scatter(
        top_gap["priority_score_0_100"],
        y_positions,
        s=marker_sizes,
        c=top_gap["mean_model_confidence"],
        cmap="viridis",
        alpha=0.82,
        edgecolors="white",
        linewidths=0.8,
    )
    ax.set_yticks(y_positions, top_gap["label"])
    for position, row in top_gap.iterrows():
        ax.annotate(
            f"{int(row['news_distinct_articles'])} articles",
            (row["priority_score_0_100"], position),
            xytext=(9, 0),
            textcoords="offset points",
            va="center",
            fontsize=7.5,
        )
    ax.set_xlabel("Gap-priority score")
    ax.set_ylabel("")
    ax.set_title("Priority candidates combine recurrence, confidence and local health links", loc="left")
    ax.grid(axis="y", visible=False)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.set_xlim(top_gap["priority_score_0_100"].min() - 1.8, top_gap["priority_score_0_100"].max() + 2.2)
    fig.colorbar(scatter_obj, ax=ax, label="Mean model confidence")
    fig.tight_layout()
    writer.save(fig, "gap_priority", "Priority candidates combine recurrence, confidence and local health links", "The score is a transparent discovery heuristic, not a risk or prevalence estimate.", top_gap)

    # 16. Test coverage classes.
    test_counts = gaps.groupby("test_coverage_class", as_index=False).agg(food_candidate_pairs=("adulterant_or_substitute", "size"), news_articles=("news_distinct_articles", "sum"))
    horizontal_bar(writer, test_counts, "test_coverage_class", "food_candidate_pairs", "test_coverage", "Candidate-specific test representation", "A missing graph edge triggers method retrieval and validation; it does not prove no test exists.", COLORS["cyan"])

    # 17. Health effects.
    news_health = [row for row in tables["news_health_edges"]]
    effect_ids: defaultdict[str, set[str]] = defaultdict(set)
    for row in news_health:
        effect_ids[row["effect"]].add(row["source_id"])
    effect_frame = pd.DataFrame([{"effect": effect, "distinct_articles": len(ids)} for effect, ids in effect_ids.items()]).nlargest(18, "distinct_articles")
    horizontal_bar(writer, effect_frame, "effect", "distinct_articles", "health_effects", "Most represented health effects in news triplets", "These are reported effects in the local corpus, not incidence estimates.", COLORS["red"])

    # 18. Effects by food pool.
    pool_effect: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    for row in news_health:
        pool_effect[(row["food_pool"], row["effect"])].add(row["source_id"])
    top_effects = [item[0] for item in Counter({effect: len({source for (pool, e), ids in pool_effect.items() if e == effect for source in ids}) for _, effect in pool_effect}).most_common(15)]
    effect_matrix = pd.DataFrame([{"food_pool": pool, "effect": effect, "articles": len(ids)} for (pool, effect), ids in pool_effect.items() if effect in top_effects]).pivot_table(index="effect", columns="food_pool", values="articles", aggfunc="sum", fill_value=0)
    fig, ax = plt.subplots(figsize=(9.8, 6.8))
    image = ax.imshow(effect_matrix.values, aspect="auto", cmap="Blues")
    ax.set_xticks(np.arange(len(effect_matrix.columns)), effect_matrix.columns)
    ax.set_yticks(np.arange(len(effect_matrix.index)), effect_matrix.index)
    ax.set_title("Health-effect representation by source food pool", loc="left")
    fig.colorbar(image, ax=ax, label="Distinct articles")
    fig.tight_layout()
    writer.save(fig, "effect_food_heatmap", "Health-effect representation by source food pool", "Pool-level differences reflect corpus composition and repeated coverage.", effect_matrix.reset_index())

    # 19. Direct causal chains.
    direct = chains[chains["same_article_direct_chain_articles"] > 0].nlargest(18, ["same_article_direct_chain_articles", "chain_support_score"]).copy()
    direct["chain"] = direct["food"] + " -> " + direct["adulterant_or_substitute"] + " -> " + direct["effect"]
    horizontal_bar(writer, direct, "chain", "same_article_direct_chain_articles", "direct_causal_chains", "Directly joined food-adulterant-effect chains", "Both graph edges occur in the same article and share the canonical adulterant node.", COLORS["red"])

    # 20. Chain classes.
    chain_classes = chains.groupby("chain_class", as_index=False).agg(chains=("effect", "size"), mean_support=("chain_support_score", "mean"))
    horizontal_bar(writer, chain_classes, "chain_class", "chains", "causal_chain_classes", "Causal-chain evidence levels", "Repository-inferred chains are hypotheses and are separated from same-article direct joins.", COLORS["purple"])

    # 21. Reverse renal-failure hypothesis.
    renal = reverse[(reverse["food"] == "Milk") & (reverse["observed_effect"] == "renal failure") & (reverse["articles_also_naming_candidate"] >= 1)].nlargest(14, "articles_also_naming_candidate")
    horizontal_bar(writer, renal, "candidate_adulterant", "observed_article_share_pct", "renal_reverse_hypothesis", "Candidate presence given Milk and renal-failure context", "The percentage is an article association share, not a diagnostic or causal probability.", COLORS["blue"], "{:.1f}%")

    # 22. Food specificity.
    specificity = pd.DataFrame(tables["food_specificity"])
    top_candidates = list(specificity.groupby("adulterant_or_substitute")["distinct_articles"].sum().nlargest(14).index)
    top_foods = list(specificity.groupby("food")["distinct_articles"].sum().nlargest(10).index)
    spec_matrix = specificity[specificity["adulterant_or_substitute"].isin(top_candidates) & specificity["food"].isin(top_foods)].pivot_table(index="adulterant_or_substitute", columns="food", values="distinct_articles", aggfunc="sum", fill_value=0)
    row_sums = spec_matrix.sum(axis=1).replace(0, 1)
    share_matrix = spec_matrix.div(row_sums, axis=0) * 100
    fig, ax = plt.subplots(figsize=(11.2, 7.0))
    image = ax.imshow(share_matrix.values, aspect="auto", cmap="PuBuGn", vmin=0, vmax=100)
    ax.set_xticks(np.arange(len(share_matrix.columns)), share_matrix.columns, rotation=35, ha="right")
    ax.set_yticks(np.arange(len(share_matrix.index)), share_matrix.index)
    ax.set_title("Food specificity of leading adulterant candidates", loc="left")
    fig.colorbar(image, ax=ax, label="Share of candidate's explicit food-link articles (%)")
    fig.tight_layout()
    writer.save(fig, "food_specificity", "Food specificity of leading adulterant candidates", "This directly answers whether a candidate is represented across foods or concentrated in one food matrix.", share_matrix.reset_index())

    # 23. Yearly food-pool article coverage for direct pairs.
    annual_pool: defaultdict[tuple[int, str], set[str]] = defaultdict(set)
    for row in news_edges:
        date = clean(tables["articles"].get(row["source_id"], {}).get("date"))
        match = re.match(r"^(\d{4})", date)
        if match:
            annual_pool[(int(match.group(1)), row["food_pool"])].add(row["source_id"])
    annual_frame = pd.DataFrame([{"year": year, "food_pool": pool, "distinct_articles": len(ids)} for (year, pool), ids in annual_pool.items()])
    recent = annual_frame[(annual_frame["year"] >= max(annual_frame["year"].max() - 12, annual_frame["year"].min())) & (annual_frame["year"] <= 2026)]
    fig, ax = plt.subplots(figsize=(10.5, 5.9))
    for pool, color in [("Oil", COLORS["gold"]), ("Ghee", COLORS["green"]), ("Milk", COLORS["blue"])]:
        series = recent[recent["food_pool"] == pool].sort_values("year")
        ax.plot(series["year"], series["distinct_articles"], marker="o", label=pool, color=color, linewidth=2)
    ax.set_ylabel("Distinct articles with direct food-adulterant triplets")
    ax.set_title("Publication-year coverage by food pool", loc="left")
    ax.legend()
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    writer.save(fig, "time_food_pool", "Publication-year coverage by food pool", "Publication date is not incident date; 2026 is partial and collection intensity varies.", recent)

    # 24. Top candidate trajectories.
    temporal = pd.DataFrame(tables["temporal"])
    top = list(temporal.groupby("adulterant_or_substitute")["distinct_articles"].sum().nlargest(6).index)
    trend = temporal[(temporal["adulterant_or_substitute"].isin(top)) & (temporal["publication_year"] >= 2014) & (temporal["publication_year"] <= 2026)].groupby(["publication_year", "adulterant_or_substitute"], as_index=False)["distinct_articles"].sum()
    fig, ax = plt.subplots(figsize=(10.7, 6.1))
    for candidate in top:
        series = trend[trend["adulterant_or_substitute"] == candidate].sort_values("publication_year")
        ax.plot(series["publication_year"], series["distinct_articles"], marker="o", linewidth=1.7, label=candidate)
    ax.set_ylabel("Distinct articles")
    ax.set_title("Publication trajectories of leading candidates", loc="left")
    ax.legend(ncol=2, fontsize=8)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    writer.save(fig, "candidate_time_series", "Publication trajectories of leading candidates", "Trends describe the collected corpus and repeated reporting, not population prevalence.", trend)

    # 25. Regulatory actions.
    actions = pd.DataFrame([row for row in tables["actions"] if row["corpus"] == "News"])
    action_top = actions.groupby(["predicate", "object"], as_index=False)["distinct_sources"].sum().nlargest(16, "distinct_sources")
    action_top["relation"] = action_top["predicate"] + " -> " + action_top["object"]
    horizontal_bar(writer, action_top, "relation", "distinct_sources", "regulatory_actions", "Most represented regulatory-action relations", "Action relations complement food chemistry by showing what authorities or operators did.", COLORS["green"])

    # 26. Detection methods.
    methods = pd.DataFrame([row for row in tables["methods"] if row["corpus"] == "News"])
    method_top = methods.groupby(["predicate", "object"], as_index=False)["distinct_sources"].sum().nlargest(16, "distinct_sources")
    method_top["relation"] = method_top["predicate"] + " -> " + method_top["object"]
    horizontal_bar(writer, method_top, "relation", "distinct_sources", "detection_methods", "Most represented news detection and method relations", "News-only method language; a method mention is not automatically a validated food-specific protocol.", COLORS["cyan"])

    # 27. Source concentration.
    publishers = pd.DataFrame(tables["publishers"])
    publisher_top = publishers.groupby("publisher", as_index=False)["distinct_articles"].sum().nlargest(18, "distinct_articles")
    horizontal_bar(writer, publisher_top, "publisher", "distinct_articles", "publisher_concentration", "News-source concentration", "Concentrated sourcing can amplify repeated incidents and editorial framing.", COLORS["gray"])

    # 28. Confidence sensitivity.
    sensitivity = pd.DataFrame(tables["confidence_sensitivity"])
    fig, ax = plt.subplots(figsize=(10.1, 5.9))
    for pool, color in [("Oil", COLORS["gold"]), ("Ghee", COLORS["green"]), ("Milk", COLORS["blue"]), ("All news", COLORS["navy"])]:
        series = sensitivity[sensitivity["food_pool"] == pool]
        ax.plot(series["minimum_confidence"], series["share_of_pool_rows_pct"], marker="o", label=pool, color=color, linewidth=2)
    ax.set_xlabel("Minimum confidence threshold")
    ax.set_ylabel("Share of rows retained (%)")
    ax.set_title("Sensitivity to confidence thresholds", loc="left")
    ax.legend()
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    writer.save(fig, "confidence_sensitivity", "Sensitivity to confidence thresholds", "Threshold views are reported alongside the all-triplet result; no historical status field is used.", sensitivity)

    # 29. Causal network.
    network_rows = direct.nlargest(12, "same_article_direct_chain_articles")
    graph = nx.DiGraph()
    for _, row in network_rows.iterrows():
        graph.add_edge(row["food"], row["adulterant_or_substitute"], weight=row["same_article_direct_chain_articles"], kind="food-adulterant")
        graph.add_edge(row["adulterant_or_substitute"], row["effect"], weight=row["same_article_direct_chain_articles"], kind="adulterant-effect")
    pos = nx.spring_layout(graph, seed=42, k=1.1)
    fig, ax = plt.subplots(figsize=(11.3, 8.0))
    node_colors = [COLORS["blue"] if node in set(network_rows["food"]) else COLORS["red"] if node in set(network_rows["effect"]) else COLORS["gold"] for node in graph.nodes]
    widths = [0.8 + math.log1p(graph.edges[edge]["weight"]) for edge in graph.edges]
    nx.draw_networkx_nodes(graph, pos, node_color=node_colors, node_size=900, alpha=0.90, ax=ax)
    nx.draw_networkx_edges(graph, pos, width=widths, edge_color=COLORS["gray"], arrows=True, arrowsize=14, ax=ax)
    nx.draw_networkx_labels(graph, pos, font_size=7.2, ax=ax)
    ax.set_title("Leading directly joined causal-chain network", loc="left")
    ax.axis("off")
    fig.tight_layout()
    writer.save(fig, "causal_network", "Leading directly joined causal-chain network", "Blue nodes are foods, gold nodes are adulterants, and red nodes are effects.", network_rows[["food", "adulterant_or_substitute", "effect", "same_article_direct_chain_articles"]])

    return writer.close()


def export_data_products(tables: dict[str, object]) -> dict[str, Path]:
    OUT.mkdir(parents=True, exist_ok=True)
    TABLES.mkdir(parents=True, exist_ok=True)
    baseline_path = OUT / f"FSSAI_Baseline_Triplets_{DATE_TAG}.csv"
    news_path = OUT / f"News_Triplets_Oil_Ghee_Milk_{DATE_TAG}.csv"
    combined_path = OUT / f"Combined_FSSAI_Baseline_and_News_Triplets_{DATE_TAG}.csv"
    provenance_path = OUT / f"Combined_Triplet_Provenance_{DATE_TAG}.csv"

    shutil.copy2(BASELINE_PATH, baseline_path)
    write_csv(news_path, tables["news"], BASE_COLUMNS)
    write_csv(combined_path, tables["baseline"] + tables["news"], BASE_COLUMNS)
    normalized_fields = list(tables["normalized"][0])
    write_csv(provenance_path, tables["normalized"], normalized_fields)

    table_map = {
        "core_profiles": tables["profiles"],
        "food_pool_profiles": tables["food_profiles"],
        "predicate_comparison": tables["predicates"],
        "food_adulterant_edges": tables["all_food_edges"],
        "food_adulterant_summary": tables["food_adulterants"],
        "gap_candidates": tables["gaps"],
        "test_priorities": tables["test_priorities"],
        "health_relations": tables["health_summary"],
        "detection_methods": tables["methods"],
        "regulatory_actions": tables["actions"],
        "direct_and_inferred_causal_chains": tables["causal_chains"],
        "contextual_causal_hypotheses": tables["contextual_chains"],
        "reverse_hypotheses": tables["reverse_hypotheses"],
        "food_specificity": tables["food_specificity"],
        "publication_time_series": tables["temporal"],
        "confidence_sensitivity": tables["confidence_sensitivity"],
        "publisher_concentration": tables["publishers"],
    }
    for name, rows in table_map.items():
        frame = pd.DataFrame(rows)
        frame.to_csv(TABLES / f"{name}.csv", index=False, encoding="utf-8-sig")
    return {
        "baseline": baseline_path,
        "news": news_path,
        "combined": combined_path,
        "provenance": provenance_path,
    }


def build_workbook(tables: dict[str, object], manifest: list[dict]) -> Path:
    workbook_path = OUT / f"FSSAI_Triplet_Intelligence_Workbook_{DATE_TAG}.xlsx"
    overlap_rows = [{"metric": key, "value": value} for key, value in tables["overlap"].items()]
    executive = [
        {"metric": "FSSAI baseline triplet rows", "value": len(tables["baseline"])},
        {"metric": "News triplet rows", "value": len(tables["news"])},
        {"metric": "Combined triplet rows", "value": len(tables["baseline"]) + len(tables["news"])},
        {"metric": "News articles represented", "value": len({row["source_id"] for row in tables["news"]})},
        {"metric": "Input news articles", "value": len(tables["articles"])},
        {"metric": "News direct food-adulterant rows", "value": len(tables["news_food_edges"])},
        {"metric": "News health-effect rows", "value": len(tables["news_health_edges"])},
        {"metric": "Food-adulterant-effect chains", "value": len(tables["causal_chains"])},
        {"metric": "Same-article direct chains", "value": sum(row["same_article_direct_chain_articles"] > 0 for row in tables["causal_chains"])},
        {"metric": "Exact normalized triplet overlap", "value": tables["overlap"]["exact_normalized_triplet_overlap"]},
        {"metric": "Historical triplet status used", "value": "No"},
        {"metric": "External/internet data used", "value": "No"},
    ]
    readme = [
        {"field": "Purpose", "value": "Compare the FSSAI baseline with all retained Oil, Ghee, and Milk news triplets; identify food-specific relations, candidate coverage gaps, test priorities and causal-chain hypotheses."},
        {"field": "Evidence boundary", "value": "Only local repository files are used. No internet or external facts are incorporated."},
        {"field": "News quality rule", "value": "All retained news triplets are treated equally. Model confidence is the only model-derived quality signal."},
        {"field": "Status rule", "value": "Historical accepted/corrected/schema-gap/unresolved labels are not used, displayed, filtered or weighted in this package."},
        {"field": "FSSAI wording", "value": "Absence means not represented in the supplied FSSAI baseline. It is not proof of absence from every current FSSAI document."},
        {"field": "Causal wording", "value": "Direct graph chains, contextual hypotheses and cross-source inferences are separated. None is epidemiological proof."},
        {"field": "Frequency wording", "value": "Triplet rows and news articles are not independent incidents or prevalence estimates."},
        {"field": "Time wording", "value": "Time series use publication dates, not incident dates; 2026 is partial."},
    ]
    data_dictionary = [
        {"field": "triplet_rows", "definition": "Number of subject-predicate-object rows."},
        {"field": "distinct_sources", "definition": "Unique FSSAI source documents or news articles, depending on corpus."},
        {"field": "confidence_weighted_source_support", "definition": "Sum of the maximum confidence per unique source for a relationship."},
        {"field": "priority_score_0_100", "definition": "Heuristic combining confidence, recurrence, local health linkage, absent FSSAI pair and absent method relation."},
        {"field": "directly joined chain", "definition": "Food-adulterant and adulterant-effect edges share the canonical adulterant and occur in the same article."},
        {"field": "repository-inferred chain", "definition": "Compatible edges share a canonical node but come from separate sources."},
        {"field": "observed_article_share_pct", "definition": "P(candidate named | food and effect context) measured over represented articles; not causal probability."},
        {"field": "food_pool", "definition": "Collection that supplied the article; not necessarily every affected food mentioned in the article."},
        {"field": "food", "definition": "Affected food extracted from the triplet and conservatively canonicalized."},
    ]

    sheet_frames = {
        "README": pd.DataFrame(readme),
        "Executive": pd.DataFrame(executive),
        "Corpus_Profiles": pd.DataFrame(tables["profiles"]),
        "Food_Pools": pd.DataFrame(tables["food_profiles"]),
        "Overlap": pd.DataFrame(overlap_rows),
        "Confidence_Sensitivity": pd.DataFrame(tables["confidence_sensitivity"]),
        "Predicates": pd.DataFrame(tables["predicates"]),
        "Food_Adulterants": pd.DataFrame(tables["food_adulterants"]),
        "Gap_Candidates": pd.DataFrame(tables["gaps"]),
        "Test_Priorities": pd.DataFrame(tables["test_priorities"]),
        "Health_Relations": pd.DataFrame(tables["health_summary"]),
        "Detection_Methods": pd.DataFrame(tables["methods"]),
        "Regulatory_Actions": pd.DataFrame(tables["actions"]),
        "Causal_Chains": pd.DataFrame(tables["causal_chains"]),
        "Context_Hypotheses": pd.DataFrame(tables["contextual_chains"]),
        "Reverse_Hypotheses": pd.DataFrame(tables["reverse_hypotheses"]),
        "Food_Specificity": pd.DataFrame(tables["food_specificity"]),
        "Time_Series": pd.DataFrame(tables["temporal"]),
        "Publishers": pd.DataFrame(tables["publishers"]),
        "Figure_Manifest": pd.DataFrame(manifest),
        "Data_Dictionary": pd.DataFrame(data_dictionary),
        "FSSAI_Baseline": pd.DataFrame(tables["baseline"])[BASE_COLUMNS],
        "News_Triplets": pd.DataFrame(tables["news"])[BASE_COLUMNS],
        "Combined_Triplets": pd.DataFrame(tables["baseline"] + tables["news"])[BASE_COLUMNS],
        "Normalized_Combined": pd.DataFrame(tables["normalized"]),
    }
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for name, frame in sheet_frames.items():
            frame.to_excel(writer, index=False, sheet_name=name)

    wb = load_workbook(workbook_path)
    header_fill = PatternFill("solid", fgColor="163A5F")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 34
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_index, cells in enumerate(ws.columns, start=1):
            sampled = [clean(cell.value) for cell in list(cells)[:300]]
            header = clean(ws.cell(1, column_index).value).casefold()
            width = min(48, max(10, max((len(value) for value in sampled), default=8) + 2))
            if any(token in header for token in ["evidence", "interpretation", "recommended", "method_examples", "article_title"]):
                width = 48
            elif "url" in header or "source_file" in header:
                width = 38
            ws.column_dimensions[get_column_letter(column_index)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_index, cell in enumerate(ws[1], start=1):
            header = clean(cell.value).casefold()
            letter = get_column_letter(column_index)
            if "confidence" in header and ws.max_row > 2:
                ws.conditional_formatting.add(
                    f"{letter}2:{letter}{ws.max_row}",
                    ColorScaleRule(start_type="min", start_color="FCE8E6", mid_type="percentile", mid_value=50, mid_color="FFF3CD", end_type="max", end_color="D9EAD3"),
                )
            if "url" in header or "source_file" in header:
                for url_cell in ws.iter_cols(min_col=column_index, max_col=column_index, min_row=2, max_row=ws.max_row):
                    for item in url_cell:
                        value = clean(item.value)
                        if value.startswith("http://") or value.startswith("https://"):
                            item.hyperlink = value
                            item.style = "Hyperlink"
    wb.save(workbook_path)
    return workbook_path


def figure_markdown(manifest: list[dict], number: int) -> str:
    item = next(row for row in manifest if row["number"] == number)
    return f"![{item['title']}]({item['png']})\n\n*Figure {number}. {item['caption']}*"


def top_food_pairs(tables: dict[str, object], target: str, limit: int = 15) -> list[dict]:
    rows = [row for row in tables["food_adulterants"] if row["corpus"] == "News" and row["specific_candidate"]]
    if target == "Oil":
        rows = [row for row in rows if "oil" in row["food"].casefold()]
    elif target == "Milk":
        rows = [row for row in rows if row["food"] in {"Milk", "Milk powder"}]
    else:
        rows = [row for row in rows if row["food"] == "Ghee"]
    return sorted(rows, key=lambda row: (-row["distinct_sources"], -float(row["mean_confidence"] or 0)))[:limit]


def build_markdown_report(tables: dict[str, object], manifest: list[dict]) -> str:
    profiles = {row["corpus"]: row for row in tables["profiles"]}
    food_profiles = {row["food_pool"]: row for row in tables["food_profiles"]}
    gaps = tables["gaps"]
    tests = tables["test_priorities"]
    chains = tables["causal_chains"]
    reverse = tables["reverse_hypotheses"]
    specificity = tables["food_specificity"]
    health = [row for row in tables["health_summary"] if row["corpus"] == "News"]
    methods = [row for row in tables["methods"] if row["corpus"] == "News"]
    actions = [row for row in tables["actions"] if row["corpus"] == "News"]
    sensitivity = tables["confidence_sensitivity"]
    overlap = tables["overlap"]

    direct_chains = [row for row in chains if row["same_article_direct_chain_articles"] > 0]
    repository_chains = [row for row in chains if row["chain_class"] == "repository-inferred from separate sources"]
    gap_class_counts = Counter(row["baseline_gap_class"] for row in gaps)
    test_class_counts = Counter(row["test_coverage_class"] for row in gaps)
    health_effect_articles: defaultdict[str, set[str]] = defaultdict(set)
    for row in tables["news_health_edges"]:
        health_effect_articles[row["effect"]].add(row["source_id"])
    top_effects = sorted(health_effect_articles.items(), key=lambda item: -len(item[1]))[:20]

    renal = next(
        (
            row
            for row in reverse
            if row["food"] == "Milk"
            and row["observed_effect"] == "renal failure"
            and row["candidate_adulterant"] == "ethylene glycol"
        ),
        None,
    )
    renal_chain = next(
        (
            row
            for row in direct_chains
            if row["food"] == "Milk"
            and row["adulterant_or_substitute"] == "ethylene glycol"
            and row["effect"] == "renal failure"
        ),
        None,
    )
    ethylene_specificity = [row for row in specificity if row["adulterant_or_substitute"] == "ethylene glycol"]
    ethylene_specificity.sort(key=lambda row: -row["distinct_articles"])

    corpus_table = markdown_table(
        ["Corpus", "Rows", "Unique triplets", "Entities", "Sources", "Mean confidence"],
        [
            [
                row["corpus"],
                f"{row['triplet_rows']:,}",
                f"{row['unique_normalized_triplets']:,}",
                f"{row['unique_entities']:,}",
                f"{row['distinct_sources']:,}",
                f"{row['mean_confidence']:.4f}",
            ]
            for row in tables["profiles"]
        ],
    )
    food_table = markdown_table(
        ["Pool", "Input articles", "Articles represented", "Triplet rows", "Mean confidence", "Rows >=0.90"],
        [
            [row["food_pool"], row["input_articles"], row["articles_with_triplets"], row["triplet_rows"], row["mean_confidence"], row["rows_confidence_ge_0_90"]]
            for row in tables["food_profiles"]
        ],
    )
    gap_table = markdown_table(
        ["Rank", "Food", "Candidate", "Articles", "Mean confidence", "FSSAI baseline class", "Method class", "Score"],
        [
            [index, row["food"], row["adulterant_or_substitute"], row["news_distinct_articles"], row["mean_model_confidence"], row["baseline_gap_class"], row["test_coverage_class"], row["priority_score_0_100"]]
            for index, row in enumerate(gaps[:25], start=1)
        ],
    )
    causal_table = markdown_table(
        ["Food", "Adulterant", "Effect", "Same-article direct chains", "FA articles", "AE articles", "Confidence", "Evidence level"],
        [
            [row["food"], row["adulterant_or_substitute"], row["effect"], row["same_article_direct_chain_articles"], row["food_adulterant_news_articles"], row["adulterant_effect_news_articles"], row["conservative_chain_confidence"], row["chain_class"]]
            for row in direct_chains[:25]
        ],
    )
    reverse_table = markdown_table(
        ["Food", "Effect context", "Candidate", "Context articles", "Candidate articles", "Observed share", "Smoothed share", "Mean confidence"],
        [
            [row["food"], row["observed_effect"], row["candidate_adulterant"], row["articles_with_food_effect_context"], row["articles_also_naming_candidate"], f"{row['observed_article_share_pct']:.1f}%", f"{row['laplace_smoothed_presence_pct']:.1f}%", row["mean_candidate_confidence"]]
            for row in [item for item in reverse if item["articles_with_food_effect_context"] >= 3 and norm(item["candidate_adulterant"]) not in GENERIC_ADULTERANTS][:30]
        ],
    )
    effect_table = markdown_table(
        ["Reported effect", "Distinct articles"],
        [[effect, len(ids)] for effect, ids in top_effects],
    )
    test_table = markdown_table(
        ["Rank", "Food", "Candidate", "Articles", "Local health basis", "Method basis", "Recommended action"],
        [
            [row["priority_rank"], row["food"], row["candidate"], row["news_distinct_articles"], row["local_health_basis"], row["method_basis_from_local_graph"], row["recommended_action"]]
            for row in tests[:20]
        ],
    )

    food_sections = []
    figure_number = {"Oil": 9, "Ghee": 10, "Milk": 11}
    for food in ["Oil", "Ghee", "Milk"]:
        profile = food_profiles[food]
        pairs = top_food_pairs(tables, food, 15)
        pair_table = markdown_table(
            ["Affected food", "Candidate", "Rows", "Articles", "Mean confidence", "Weighted support"],
            [[row["food"], row["adulterant_or_substitute"], row["triplet_rows"], row["distinct_sources"], row["mean_confidence"], row["confidence_weighted_source_support"]] for row in pairs],
        )
        food_gaps = [row for row in gaps if (food == "Oil" and "oil" in row["food"].casefold()) or (food == "Milk" and row["food"] in {"Milk", "Milk powder"}) or (food == "Ghee" and row["food"] == "Ghee")]
        gap_summary = "; ".join(f"{row['food']} -> {row['adulterant_or_substitute']} ({row['news_distinct_articles']} articles, confidence {row['mean_model_confidence']})" for row in food_gaps[:6])
        food_sections.append(
            f"""## {food}: food-specific news findings

The {food} pool contains **{profile['triplet_rows']:,} triplet rows** from **{profile['articles_with_triplets']:,} represented articles** out of {profile['input_articles']:,} inputs. Its mean confidence is {profile['mean_confidence']:.4f}. The table below uses the food explicitly named in each triplet; it does not relabel every relation according to the collection name.

{pair_table}

Leading candidate baseline gaps for this food scope are: {gap_summary or 'none among the highest-ranked local candidates'}.

{figure_markdown(manifest, figure_number[food])}
"""
        )

    report = f"""# Food-Specific Triplet Intelligence from the FSSAI Baseline and News

## Oil, Ghee and Milk | Local-data-only research report | 22 August 2026

## Abstract

This report compares **{profiles['FSSAI baseline']['triplet_rows']:,} FSSAI baseline triplets** with **{profiles['News']['triplet_rows']:,} triplets extracted from Oil, Ghee and Milk news articles**. All retained news triplets are included without using historical adjudication labels; model confidence is the only model-derived quality variable. The analysis asks four questions: what structured relations dominate official and incident-oriented knowledge, which food-adulterant and food-test relationships are not represented in the FSSAI baseline, what health and enforcement signals occur in news, and whether compatible graph paths support food-adulterant-effect causal hypotheses.

The graph contains **{len(tables['news_food_edges']):,} direct news food-adulterant/substitute rows**, **{len(tables['news_health_edges']):,} news effect rows**, and **{len(direct_chains):,} food-adulterant-effect combinations with at least one directly joined same-article chain**. The strongest requested example is Milk -> ethylene glycol -> renal failure: {renal_chain['same_article_direct_chain_articles'] if renal_chain else 0} articles contain the directly joinable chain. In the reverse article-context analysis, {renal['articles_also_naming_candidate'] if renal else 0} of {renal['articles_with_food_effect_context'] if renal else 0} Milk articles carrying renal-failure context and at least one explicit food-adulterant relation also name ethylene glycol ({renal['observed_article_share_pct'] if renal else 0:.1f}%; Laplace-smoothed {renal['laplace_smoothed_presence_pct'] if renal else 0:.1f}%). This is useful for hypothesis ranking but is not a calibrated diagnostic probability and may reflect repeated coverage of one incident.

## Evidence boundary

No internet, external database, chemical handbook, clinical reference, regulatory webpage or unstored source is used. Every descriptive claim comes from the supplied CSV/JSONL files or from deterministic joins over those files. Recommendations about validation design are methodological proposals. A missing relationship means **not represented in the supplied FSSAI baseline**, not proven absent from all current FSSAI documents.

## Executive findings

1. The combined corpus contains **{profiles['FSSAI baseline']['triplet_rows'] + profiles['News']['triplet_rows']:,} rows**, but only **{overlap['exact_normalized_triplet_overlap']} exact normalized triplets** overlap. The main reason is functional: the FSSAI baseline emphasizes definitions, ingredients, functions, standards, limits and obligations, while news emphasizes adulterants, health effects, actors, places, methods and actions.
2. News contributes **{overlap['news_only_normalized_triplets']:,} normalized triplets** and **{overlap['news_only_entities']:,} entities not exactly represented in the baseline. These are discovery leads, not automatically validated additions.
3. Leading food-specific candidate gaps include {', '.join(f"{row['food']} -> {row['adulterant_or_substitute']}" for row in gaps[:8])}. Each is ranked using confidence, unique-article recurrence, locally linked health evidence, baseline pair coverage and method coverage.
4. The graph supports explicit causal-chain analysis. Milk -> ethylene glycol -> renal failure is the dominant directly joined chain, followed by Milk -> ethylene glycol -> multi-organ failure and Milk -> ethylene glycol -> anuria.
5. Candidate food specificity is material. Ethylene glycol is concentrated in {', '.join(f"{row['food']} ({row['distinct_articles']} articles; {row['share_of_candidate_food_links_pct']:.1f}%)" for row in ethylene_specificity[:5])}. It should not be assigned equal prior relevance to Oil, Ghee and Milk.
6. Confidence is high and compressed: news mean {profiles['News']['mean_confidence']:.4f}, median {profiles['News']['median_confidence']:.2f}. Therefore confidence supports ranking but cannot substitute for entity resolution, incident deduplication, laboratory confirmation or causal validation.

## 1. Research questions

The analysis is designed around structured questions rather than a list of chemical names:

- Which subject-predicate-object relations are shared or unique across the FSSAI baseline and news?
- Which candidate adulterants are connected to which foods, and how often across distinct articles?
- Does the FSSAI baseline represent the exact food-adulterant pair, only the candidate in another role, or neither?
- Does either graph represent a candidate-specific detection method?
- Which effects are directly attributed to adulterants, and which are only contextually associated within an article?
- Given a food and observed effect, which candidate adulterants are most frequently co-represented?
- Is a candidate food-specific or spread across several food matrices?
- Which graph gaps should be routed to document verification, method validation or targeted surveillance?

## 2. Data inventory

{corpus_table}

{food_table}

The news export represents {sum(row['articles_with_triplets'] for row in tables['food_profiles']):,} article-food records across the three pools. Articles with no retained triplets remain part of corpus-coverage accounting but cannot contribute edges. The FSSAI baseline has 27 source IDs and is treated as a verified reference graph, not as a probabilistic news model output.

{figure_markdown(manifest, 1)}

{figure_markdown(manifest, 2)}

## 3. Common schema and provenance

The new FSSAI baseline, News and Combined CSVs use the same 14 columns: snippet and source identifiers, source file/type, chunk index, subject/type/ID, predicate, object/type/ID, confidence and evidence span. The news `source_file` field contains the article URL. A companion provenance file adds corpus, food pool, title, date, publisher, canonical fields and a stable triplet ID. Raw values are retained; normalization never overwrites source assertions.

Historical triplet-status labels are absent from the new CSVs, workbook and analysis. Every retained news row participates. Rejected candidates are not reconstructed because they are not present in the retained input export.

## 4. Confidence policy and sensitivity

All {profiles['News']['triplet_rows']:,} news rows and all {profiles['FSSAI baseline']['triplet_rows']:,} FSSAI baseline rows carry a finite confidence value. News values range from {profiles['News']['minimum_confidence']:.2f} to {profiles['News']['maximum_confidence']:.2f}. Counts are always reported independently from confidence. A confidence-weighted source-support measure sums the maximum confidence per source so repeated rows from one article do not contribute linearly without bound.

Confidence is not assumed calibrated across predicates, foods or models. Sensitivity tables show all rows and minimum thresholds of 0.70, 0.80, 0.90 and 0.95. The main report uses all triplets as requested.

{figure_markdown(manifest, 3)}

{figure_markdown(manifest, 4)}

{figure_markdown(manifest, 28)}

## 5. Combined knowledge-graph landscape

The FSSAI baseline contains {overlap['baseline_only_normalized_triplets']:,} normalized triples not found exactly in news; news contains {overlap['news_only_normalized_triplets']:,} not found exactly in the baseline. Only {overlap['entity_overlap']:,} normalized entity strings overlap. Exact comparison is deliberately conservative and exposes the need for identifiers: all source rows leave subject and object IDs blank, so synonym resolution currently depends on text normalization and audited alias maps.

{figure_markdown(manifest, 6)}

{figure_markdown(manifest, 7)}

The sparse overlap does not mean the FSSAI baseline is irrelevant to news. It means the graphs often encode different sides of the same problem. The baseline may define a food, limit or method while news supplies an incident, adulterant, effect, actor or action. The useful unit of gap analysis is therefore **food + candidate + role + method + effect + evidence**, not term presence alone.

## 6. News-only triplet landscape

News contains {len(tables['news_food_edges']):,} direct `hasAdulterant` or `isSubstituteFor` rows. The analysis collapses exact duplicate mentions at several levels: row count, unique canonical relation, unique article and confidence-weighted article support. Article counts remain reporting counts and may still describe repeated coverage of one incident.

{figure_markdown(manifest, 5)}

{figure_markdown(manifest, 8)}

{figure_markdown(manifest, 12)}

{figure_markdown(manifest, 13)}

{''.join(food_sections)}

## 10. Candidate gaps relative to the FSSAI baseline

The gap taxonomy separates exact food-pair coverage from contextual term coverage. Across {len(gaps):,} specific news food-candidate combinations, the classes are: {', '.join(f"{name}: {count}" for name, count in gap_class_counts.most_common())}. A candidate may therefore be present in the FSSAI baseline but absent as an adulterant of the food reported in news. This distinction prevents false conclusions from simple keyword lookup.

{gap_table}

{figure_markdown(manifest, 14)}

{figure_markdown(manifest, 15)}

The priority score is a discovery heuristic: 35% mean confidence, 25% log-scaled recurrence, 20% locally represented candidate-effect evidence, 10% absent exact food-pair relation and 10% absent candidate-method relation. It is not a national risk score. A highly ranked result should trigger source inspection and FSSAI-document verification before policy use.

## 11. Test and surveillance priorities

Candidate-method representation classes are: {', '.join(f"{name}: {count}" for name, count in test_class_counts.most_common())}. Method edges can be missing because the supplied baseline is incomplete, because the ontology lacks a link, because extraction missed it, or because the documents do not specify it. The report therefore proposes a decision workflow rather than asserting a specific laboratory protocol where none is locally represented.

{test_table}

{figure_markdown(manifest, 16)}

{figure_markdown(manifest, 26)}

For a candidate with an FSSAI baseline method relation, the next step is to retrieve and validate that method for the exact food matrix and enforcement use. For a method mentioned only in news, the article becomes a retrieval lead, not a protocol. Where neither corpus represents a method, the defensible conclusion is a **method-discovery priority**. Minimum validation metadata should include food matrix, analyte identity, sampling, blanks, spikes, recovery, precision, detection and quantification limits, confirmatory rules and chain of custody.

## 12. Health effects and reported impact

The news graph contains {len(tables['news_health_edges']):,} effect rows. The most represented effects by distinct article are:

{effect_table}

{figure_markdown(manifest, 17)}

{figure_markdown(manifest, 18)}

These relations describe what the local articles assert. They do not establish clinical incidence, attributable risk or national burden. Effects linked directly from a specific adulterant are stronger graph evidence than an effect connected only to a food, event or generic adulteration statement.

## 13. Causal-chain construction

A food-adulterant-effect chain requires two compatible edges:

1. Food -> `hasAdulterant` -> candidate, or candidate -> `isSubstituteFor` -> food.
2. The same canonical candidate -> an effect relation -> health effect.

Three evidence levels are kept separate. **Direct same-article chains** share a canonical candidate and occur in one article. **Contextual hypotheses** pair a food-adulterant relation with a health effect in the same article without an explicit candidate-effect edge. **Repository-inferred chains** join compatible direct edges from separate sources. Chain confidence is the lower of the mean edge confidences, and support is increased by independent source representation. No chain is called epidemiological proof.

{causal_table}

{figure_markdown(manifest, 19)}

{figure_markdown(manifest, 20)}

{figure_markdown(manifest, 29)}

The workbook retains example URLs and both evidence spans for every direct chain. This makes each path auditable and permits human rejection of bad canonical joins.

## 14. Milk, ethylene glycol and renal failure

The local graph directly supports the example raised for this report. Milk -> ethylene glycol appears in {renal_chain['food_adulterant_news_articles'] if renal_chain else 0} distinct news articles; ethylene glycol -> renal failure appears in {renal_chain['adulterant_effect_news_articles'] if renal_chain else 0}; and {renal_chain['same_article_direct_chain_articles'] if renal_chain else 0} articles contain both joinable edges. The conservative chain confidence is {renal_chain['conservative_chain_confidence'] if renal_chain else 0}.

Example food-adulterant evidence: “{clean(renal_chain['example_food_adulterant_evidence']) if renal_chain else ''}”

Example adulterant-effect evidence: “{clean(renal_chain['example_adulterant_effect_evidence']) if renal_chain else ''}”

Example source: {renal_chain['example_article_url'] if renal_chain else ''}

Repeated articles may cover the same underlying incident. The result establishes strong repository representation of the chain, not {renal_chain['same_article_direct_chain_articles'] if renal_chain else 0} independent toxicological confirmations.

## 15. Reverse inference: candidate given food and effect

For a represented food F, effect E and candidate A, the report calculates:

`observed share = articles containing F, E and A / articles containing F and E with at least one food-adulterant relation`

It also reports a binary Laplace-smoothed share `(n + 1) / (N + 2)`. Multiple candidates may occur in one article, so candidate shares need not sum to 100%. This is a retrieval and prioritization score, not the probability that a patient or sample was exposed to the candidate.

{reverse_table}

{figure_markdown(manifest, 21)}

For Milk and renal failure, ethylene glycol is named in {renal['articles_also_naming_candidate'] if renal else 0}/{renal['articles_with_food_effect_context'] if renal else 0} represented articles ({renal['observed_article_share_pct'] if renal else 0:.1f}%). This is substantially more represented than urea or detergent in the same context. The correct operational use is to prioritize ethylene-glycol evidence retrieval and testing consideration while keeping alternatives open.

## 16. Food specificity

Food specificity asks whether a candidate's direct food-link articles concentrate in one matrix. It is based on explicitly extracted food-candidate relations, not the article pool label. This matters because an effect associated with a candidate in Milk should not automatically imply equal relevance in Oil or Ghee.

{figure_markdown(manifest, 22)}

For ethylene glycol, the represented distribution is: {', '.join(f"{row['food']}: {row['distinct_articles']} articles ({row['share_of_candidate_food_links_pct']:.1f}%)" for row in ethylene_specificity[:10])}. Small denominators and generic food labels remain visible in the workbook.

## 17. Time series

Time-series tables use publication year. They count unique articles with direct food-adulterant relations and are not adjusted for crawler intensity, duplicate reporting or partial-year coverage. They are appropriate for corpus monitoring and emergence detection, not prevalence estimation.

{figure_markdown(manifest, 23)}

{figure_markdown(manifest, 24)}

## 18. Regulatory actions and response intelligence

News adds actions, bodies, targets, operators and incident findings that are sparse in standards-oriented text. These relations are relevant to an operational food-safety graph even when they do not change the food chemistry.

{figure_markdown(manifest, 25)}

The action table in the workbook preserves subject, predicate, object, source count, confidence and example evidence. Repeated action reports should be clustered by incident before counting enforcement events.

## 19. Source concentration and duplication

The analysis uses unique articles rather than rows for recurrence, but syndicated and follow-up coverage can still describe one event. Publisher concentration and known title clusters show why article frequency cannot be read as incident frequency.

{figure_markdown(manifest, 27)}

The causal and reverse-inference results should ultimately be repeated after event-level deduplication. Until then, article count measures the strength of corpus representation, not independent-event count.

## 20. Interpreting “FSSAI gaps”

Four different phenomena can produce a missing edge:

1. **Document coverage candidate:** the supplied baseline does not contain the concept or relation.
2. **Extraction gap:** a source document may contain it, but the verified triplet file does not.
3. **Terminology gap:** baseline and news use different labels for the same entity.
4. **Ontology gap:** both concepts exist, but the relation needed to connect food, candidate, method, effect or action is absent.

The report can discover and rank these cases, but only source-document inspection distinguishes them. Accordingly, all gap labels say “not represented in the FSSAI baseline.”

## 21. AI-for-social-good research design

A defensible system should retain seven connected layers: immutable source text; atomic triplets with evidence and confidence; conservative entity resolution; article and incident deduplication; causal-path construction with evidence levels; FSSAI baseline alignment; and human/laboratory verification queues. The output is a decision-support graph, not an autonomous regulatory conclusion.

The key evaluation units are food-candidate pair precision, candidate-effect edge precision, entity-link precision, article-to-event cluster precision, method-link completeness, evidence-span exactness and calibration of confidence against human review. Reverse hypotheses should be evaluated as retrieval rankings using held-out, event-deduplicated cases.

## 22. Limitations

- News coverage is selective, duplicated and editorially mediated.
- Article count is not incident count; incident count is not prevalence.
- Confidence is model-supplied and not demonstrated to be calibrated.
- Blank subject/object identifiers force lexical canonicalization.
- Alias normalization can create false joins; raw surfaces and evidence are retained for review.
- Same-article context is not a causal assertion.
- Repository-inferred chains join claims across sources and are hypotheses only.
- Publication date is not incident date.
- The FSSAI baseline is a supplied verified extract, not a guaranteed exhaustive representation of every current FSSAI document.
- No external toxicology or analytical-method source is used; health and method conclusions are bounded by local graph content.

## 23. Conclusions

The combined graph is most useful as a structured gap-discovery and hypothesis-routing system. News contributes food-specific adulterants, effects, detection leads, actions and event context that the FSSAI baseline often does not connect. The strongest locally represented causal pattern is Milk -> ethylene glycol -> renal failure, and food-specificity analysis shows why that evidence should not be transferred indiscriminately to Oil or Ghee.

The practical next step is not automatic ontology expansion. It is a ranked verification program: inspect the source article and FSSAI documents; validate entity identity and food role; deduplicate incidents; retrieve or develop a matrix-appropriate method; and record the resulting accepted, rejected or newly modeled edge. The package supplies all rows, evidence, confidence, rankings and figure data required for that process.

## Appendix A. Top 25 candidate gaps

{gap_table}

## Appendix B. Top 25 direct causal chains

{causal_table}

## Appendix C. Reverse hypotheses with at least three context articles

{reverse_table}

## Appendix D. Reproducibility

Inputs are `triplets_verified.csv`, `News_Triplets_Oil_Ghee_Milk_20260820.csv`, and the three local article JSONL files listed in `package_manifest.json`. Run `python build_triplet_intelligence_report.py` from the repository root. Every output has a SHA-256 digest in the package manifest. No network call is made by the builder.
"""
    report_path = OUT / f"FSSAI_Triplet_Intelligence_Report_{DATE_TAG}.md"
    report_path.write_text(report.strip() + "\n", encoding="utf-8")
    return report


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def strip_markdown(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", value)
    return value.replace("**", "").replace("`", "").replace("*", "")


def add_docx_table(document: Document, lines: list[str]) -> None:
    parsed = [[strip_markdown(cell.strip()) for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(parsed) < 3:
        return
    table = document.add_table(rows=1, cols=len(parsed[0]))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for index, value in enumerate(parsed[0]):
        cell = table.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, "163A5F")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(7)
    for row_values in parsed[2:]:
        cells = table.add_row().cells
        for index, value in enumerate(row_values[: len(cells)]):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(6.7)
    document.add_paragraph()


def markdown_to_docx(markdown_text: str, output_path: Path | None = None) -> Path:
    path = output_path or OUT / f"FSSAI_Triplet_Intelligence_Report_{DATE_TAG}.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)
    document.styles["Normal"].paragraph_format.space_after = Pt(4)
    for name, size, color in [
        ("Title", 24, RGBColor(22, 58, 95)),
        ("Heading 1", 16, RGBColor(22, 58, 95)),
        ("Heading 2", 12.5, RGBColor(43, 106, 159)),
        ("Heading 3", 10.5, RGBColor(59, 122, 87)),
    ]:
        document.styles[name].font.name = "Aptos Display"
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.color.rgb = color
        document.styles[name].font.bold = True
    header = section.header.paragraphs[0]
    header.text = "FSSAI baseline and Oil/Ghee/Milk news triplet intelligence | Local data only"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(95, 105, 115)
    add_page_number(section.footer.paragraphs[0])
    document.core_properties.title = "Food-Specific Triplet Intelligence from the FSSAI Baseline and News"
    document.core_properties.subject = "Oil, Ghee and Milk triplet analytics, baseline gaps and causal-chain hypotheses"
    document.core_properties.author = "Local reproducible analysis"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Food-Specific Triplet Intelligence\nfrom the FSSAI Baseline and News")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Oil, Ghee and Milk | Local-data-only research report\n22 August 2026")
    run.font.size = Pt(12.5)
    run.font.color.rgb = RGBColor(59, 122, 87)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run("No internet or external evidence used. Baseline absence is not proof of absence from all current FSSAI documents.").bold = True
    document.add_page_break()

    lines = markdown_text.splitlines()
    index = 0
    skipped_title = False
    paragraph_buffer: list[str] = []

    def flush() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            document.add_paragraph(strip_markdown(" ".join(paragraph_buffer)))
            paragraph_buffer = []

    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            flush()
            index += 1
            continue
        if stripped.startswith("!["):
            flush()
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
            if match:
                image_path = OUT / match.group(2)
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.65))
            index += 1
            continue
        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            heading = strip_markdown(stripped[level:].strip())
            if level == 1 and not skipped_title:
                skipped_title = True
            else:
                document.add_heading(heading, level=min(level, 3))
            index += 1
            continue
        if stripped.startswith("|"):
            flush()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_docx_table(document, table_lines)
            continue
        if re.match(r"^[-*] ", stripped):
            flush()
            document.add_paragraph(strip_markdown(stripped[2:]), style="List Bullet")
            index += 1
            continue
        if re.match(r"^\d+\. ", stripped):
            flush()
            document.add_paragraph(strip_markdown(re.sub(r"^\d+\. ", "", stripped)), style="List Number")
            index += 1
            continue
        paragraph_buffer.append(stripped)
        index += 1
    flush()
    document.save(path)
    return path


def csv_header(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return next(csv.reader(handle))


def csv_data_rows(path: Path) -> int:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        next(reader)
        return sum(1 for _ in reader)


def validate_release(
    tables: dict[str, object],
    data_paths: dict[str, Path],
    figure_manifest: list[dict],
    workbook_path: Path,
    report_path: Path,
    docx_path: Path,
) -> dict[str, object]:
    errors: list[str] = []
    checks: list[dict[str, object]] = []

    def check(name: str, condition: bool, detail: str) -> None:
        checks.append({"check": name, "passed": bool(condition), "detail": detail})
        if not condition:
            errors.append(f"{name}: {detail}")

    expected_counts = {
        "baseline": len(tables["baseline"]),
        "news": len(tables["news"]),
        "combined": len(tables["baseline"]) + len(tables["news"]),
    }
    for key in ["baseline", "news", "combined"]:
        path = data_paths[key]
        header = csv_header(path)
        rows = csv_data_rows(path)
        check(
            f"{key} compatible CSV schema",
            header == BASE_COLUMNS,
            f"expected {BASE_COLUMNS}; found {header}",
        )
        check(
            f"{key} compatible CSV row count",
            rows == expected_counts[key],
            f"expected {expected_counts[key]}; found {rows}",
        )
        check(
            f"{key} compatible CSV excludes historical status",
            "status" not in {column.casefold() for column in header},
            f"header={header}",
        )

    check(
        "FSSAI baseline export is byte-identical",
        sha256(data_paths["baseline"]) == sha256(BASELINE_PATH),
        f"source={sha256(BASELINE_PATH)} export={sha256(data_paths['baseline'])}",
    )
    check(
        "normalized provenance row count",
        csv_data_rows(data_paths["provenance"]) == expected_counts["combined"],
        f"expected {expected_counts['combined']}",
    )
    normalized_ids = [row["triplet_id"] for row in tables["normalized"]]
    check(
        "stable normalized triplet IDs are unique",
        len(normalized_ids) == len(set(normalized_ids)),
        f"rows={len(normalized_ids)} unique={len(set(normalized_ids))}",
    )
    news_confidence = [float(row["confidence"]) for row in tables["news"]]
    check(
        "all news model confidences are finite",
        len(news_confidence) == len(tables["news"]) and all(math.isfinite(value) for value in news_confidence),
        f"finite={sum(math.isfinite(value) for value in news_confidence)} rows={len(tables['news'])}",
    )
    check(
        "all news source links are populated",
        all(clean(row["source_file"]) for row in tables["news"]),
        "news source_file must retain an article link or source locator",
    )

    check(
        "figure inventory",
        len(figure_manifest) == 29,
        f"expected 29; found {len(figure_manifest)}",
    )
    missing_figure_files = []
    for item in figure_manifest:
        for field in ["png", "data_csv"]:
            candidate = OUT / item[field]
            if not candidate.exists() or candidate.stat().st_size == 0:
                missing_figure_files.append(str(candidate))
    check(
        "all figure and figure-data files exist",
        not missing_figure_files,
        "; ".join(missing_figure_files) or "all files present",
    )

    renal_chain = next(
        (
            row
            for row in tables["causal_chains"]
            if row["food"] == "Milk"
            and row["adulterant_or_substitute"] == "ethylene glycol"
            and row["effect"] == "renal failure"
        ),
        None,
    )
    renal_reverse = next(
        (
            row
            for row in tables["reverse_hypotheses"]
            if row["food"] == "Milk"
            and row["observed_effect"] == "renal failure"
            and row["candidate_adulterant"] == "ethylene glycol"
        ),
        None,
    )
    check(
        "Milk-ethylene glycol-renal failure direct chain",
        renal_chain is not None and renal_chain["same_article_direct_chain_articles"] == 23,
        f"found={renal_chain}",
    )
    check(
        "Milk renal-failure reverse hypothesis denominator",
        renal_reverse is not None
        and renal_reverse["articles_also_naming_candidate"] == 23
        and renal_reverse["articles_with_food_effect_context"] == 41,
        f"found={renal_reverse}",
    )

    report_text = report_path.read_text(encoding="utf-8")
    check(
        "report uses requested FSSAI baseline terminology",
        "strict csv" not in report_text.casefold() and "strict baseline" not in report_text.casefold(),
        "report must call the reference corpus the FSSAI baseline",
    )
    check(
        "report states local-data-only boundary",
        "No internet" in report_text and "local" in report_text.casefold(),
        "missing explicit source boundary",
    )
    check(
        "workbook is readable",
        workbook_path.exists() and len(load_workbook(workbook_path, read_only=True).sheetnames) >= 20,
        f"path={workbook_path}",
    )
    check(
        "DOCX report exists",
        docx_path.exists() and docx_path.stat().st_size > 100_000,
        f"path={docx_path} bytes={docx_path.stat().st_size if docx_path.exists() else 0}",
    )

    validation = {
        "passed": not errors,
        "checks": checks,
        "errors": errors,
        "metrics": {
            "fssai_baseline_rows": len(tables["baseline"]),
            "news_rows": len(tables["news"]),
            "combined_rows": len(tables["baseline"]) + len(tables["news"]),
            "news_articles_represented": len({row["source_id"] for row in tables["news"]}),
            "direct_news_food_adulterant_rows": len(tables["news_food_edges"]),
            "news_health_effect_rows": len(tables["news_health_edges"]),
            "figures": len(figure_manifest),
        },
    }
    path = OUT / "release_validation.json"
    path.write_text(json.dumps(validation, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    if errors:
        raise RuntimeError("Release validation failed:\n- " + "\n- ".join(errors))
    return validation


def build_package_manifest(
    tables: dict[str, object],
    data_paths: dict[str, Path],
    figure_manifest: list[dict],
    validation: dict[str, object],
) -> Path:
    input_paths = [BASELINE_PATH, NEWS_PATH, *ARTICLE_PATHS]
    inputs = [
        {
            "path": str(path.relative_to(ROOT)),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        }
        for path in input_paths
    ]
    output_files = sorted(
        path
        for path in OUT.rglob("*")
        if path.is_file() and path.name not in {"package_manifest.json", "SHA256SUMS.txt"}
    )
    outputs = [
        {
            "path": str(path.relative_to(OUT)).replace("\\", "/"),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        }
        for path in output_files
    ]
    manifest = {
        "package": OUT.name,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "analysis_contract": {
            "reference_name": "FSSAI baseline",
            "foods": ["Oil", "Ghee", "Milk"],
            "historical_news_status_used": False,
            "news_model_confidence_used": True,
            "internet_or_external_data_used": False,
            "compatible_triplet_columns": BASE_COLUMNS,
            "causal_claim_boundary": "article-level graph association and hypothesis ranking; not epidemiological proof or calibrated diagnosis",
            "baseline_absence_boundary": "not represented in the supplied FSSAI baseline; not proof of absence from all FSSAI documents",
        },
        "metrics": validation["metrics"],
        "figures": len(figure_manifest),
        "inputs": inputs,
        "outputs": outputs,
        "validation": {"passed": validation["passed"], "checks": len(validation["checks"])},
    }
    path = OUT / "package_manifest.json"
    path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    checksums = sorted(
        candidate
        for candidate in OUT.rglob("*")
        if candidate.is_file() and candidate.name != "SHA256SUMS.txt"
    )
    (OUT / "SHA256SUMS.txt").write_text(
        "\n".join(
            f"{sha256(candidate)}  {str(candidate.relative_to(OUT)).replace(chr(92), '/')}"
            for candidate in checksums
        )
        + "\n",
        encoding="ascii",
    )
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    TABLES.mkdir(parents=True, exist_ok=True)
    print("Building normalized analytical tables...")
    tables = build_all_tables()
    print("Writing compatible CSV exports and analytical tables...")
    data_paths = export_data_products(tables)
    print("Rendering figures and figure-data exports...")
    figure_manifest = generate_figures(tables)
    print("Building analytical workbook...")
    workbook_path = build_workbook(tables, figure_manifest)
    print("Writing Markdown and DOCX reports...")
    report_text = build_markdown_report(tables, figure_manifest)
    report_path = OUT / f"FSSAI_Triplet_Intelligence_Report_{DATE_TAG}.md"
    docx_path = markdown_to_docx(report_text)
    print("Validating release contract...")
    validation = validate_release(
        tables,
        data_paths,
        figure_manifest,
        workbook_path,
        report_path,
        docx_path,
    )
    manifest_path = build_package_manifest(tables, data_paths, figure_manifest, validation)
    print(
        json.dumps(
            {
                "output_directory": str(OUT),
                "fssai_baseline_rows": len(tables["baseline"]),
                "news_rows": len(tables["news"]),
                "combined_rows": len(tables["baseline"]) + len(tables["news"]),
                "figures": len(figure_manifest),
                "workbook": str(workbook_path),
                "report_markdown": str(report_path),
                "report_docx": str(docx_path),
                "manifest": str(manifest_path),
                "validation_passed": validation["passed"],
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
